"""Shared helpers: config, workbook access by column name."""
import os
import re
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK = BASE_DIR / "UH_Homes_Selections_Tracker.xlsx"


def load_env(path=BASE_DIR / ".env"):
    """Minimal .env loader (KEY=value lines) so no extra package is needed."""
    if not Path(path).exists():
        return
    for line in Path(path).read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


load_env()
WORKBOOK = Path(os.getenv("UH_WORKBOOK", str(WORKBOOK)))
if not WORKBOOK.is_absolute():
    WORKBOOK = BASE_DIR / WORKBOOK


def clean(v):
    if v is None:
        return ""
    return str(v).strip()


def norm_model(s):
    """'K-596-VS' -> 'k596vs' so formatting differences don't break matching."""
    return re.sub(r"[^a-z0-9]", "", clean(s).lower())


class Sheet:
    """Access a sheet's rows as dicts keyed by header, and write back by header."""

    def __init__(self, ws):
        self.ws = ws
        self.cols = {clean(c.value): c.column for c in ws[1] if clean(c.value)}

    def rows(self):
        for r in range(2, self.ws.max_row + 1):
            data = {h: self.ws.cell(row=r, column=c).value for h, c in self.cols.items()}
            if any(clean(v) for k, v in data.items() if k not in AUTO_OR_FORMULA):
                yield r, data

    def set(self, row, header, value):
        self.ws.cell(row=row, column=self.cols[header], value=value)


# Columns ignored when deciding whether a row is "empty"
AUTO_OR_FORMULA = {"Total Items", "Links Found", "Verified", "Needs Review",
                   "Lookup Status", "Include in Lookbook"}


def open_workbook(data_only=False):
    if not WORKBOOK.exists():
        raise SystemExit(f"Workbook not found: {WORKBOOK}")
    return load_workbook(WORKBOOK, data_only=data_only)


def manufacturer_domains(wb):
    out = {}
    for _, d in Sheet(wb["Manufacturers"]).rows():
        name, dom = clean(d.get("Manufacturer")), clean(d.get("Official Domain"))
        if name and dom:
            dom = re.sub(r"^https?://", "", dom).split("/")[0].removeprefix("www.")
            out[name.lower()] = dom.lower()
    return out


def section_order(wb):
    return [clean(r[0].value) for r in wb["Lists"].iter_rows(min_row=2, max_col=1) if clean(r[0].value)]
