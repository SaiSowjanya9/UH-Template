import io
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import app
import common
from auto_lookup import AutoLookup
from workbook_store import WorkbookStore


class CustomFieldTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.path = Path(self.folder.name) / "master.xlsx"
        shutil.copy2(common.WORKBOOK, self.path)
        self.store = WorkbookStore(self.path)
        self.patches = [patch.object(app, "store", self.store), patch.object(app, "automation", None)]
        for item in self.patches:
            item.start()
        self.client = app.app.test_client()
        self.headers = {"X-UH-Token": app.TOKEN}

    def tearDown(self):
        for item in reversed(self.patches):
            item.stop()
        self.folder.cleanup()

    def state(self):
        response = self.client.get("/api/state")
        self.assertEqual(response.status_code, 200, response.json)
        return response.json

    def post(self, path, data):
        return self.client.post(path, json={"revision": self.state()["revision"], **data}, headers=self.headers)

    def test_project_fields_persist_and_are_isolated(self):
        data = self.state()
        project = data["projects"][0]
        fields = [{"name": "Lot number", "value": "24"}, {"name": "Construction manager", "value": "Test Manager"}]
        response = self.post("/api/projects", {"values": {"Project ID": project["Project ID"]}, "custom_fields": fields})
        self.assertEqual(response.status_code, 200, response.json)
        data = self.state()
        self.assertEqual(data["projects"][0]["custom_fields"], fields)
        self.assertEqual(data["projects"][1]["custom_fields"], [])
        with patch.object(app, "store", WorkbookStore(self.path)):
            self.assertEqual(self.state()["projects"][0]["custom_fields"], fields)

    def test_selection_fields_survive_reordering_and_roundtrip(self):
        row = self.state()["selections"][0]
        fields = [{"name": "Supplier", "value": "Test supplier"}, {"name": "Budget", "value": "0"}]
        response = self.post("/api/selections", {"row": row["_row"], "values": {}, "custom_fields": fields})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.state()["selections"][0]["custom_fields"], fields)
        self.assertEqual(self.state()["selections"][1]["custom_fields"], [])
        wb, revision = self.store.snapshot()
        wb["Selections"].insert_rows(2)
        self.store.save(wb, revision)
        self.assertEqual(self.state()["selections"][0]["custom_fields"], fields)
        self.assertEqual(self.state()["selections"][0]["_row"], row["_row"] + 1)
        export = self.client.get("/api/workbook")
        response = self.client.post("/api/workbook/import", data={"file": (io.BytesIO(export.data), "tracker.xlsx"), "revision": self.state()["revision"], "confirm": "replace-with-backup"}, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.state()["selections"][0]["custom_fields"], fields)

    def test_edit_remove_and_omitted_fields(self):
        row = self.state()["selections"][0]["_row"]
        fields = [{"name": "Supplier", "value": "First"}]
        self.post("/api/selections", {"row": row, "values": {}, "custom_fields": fields})
        self.post("/api/selections", {"row": row, "values": {"Qty": "2"}})
        self.assertEqual(self.state()["selections"][0]["custom_fields"], fields)
        updated = [{"name": "Installer", "value": "Second"}]
        self.post("/api/selections", {"row": row, "values": {}, "custom_fields": updated})
        self.assertEqual(self.state()["selections"][0]["custom_fields"], updated)
        self.post("/api/selections", {"row": row, "values": {}, "custom_fields": []})
        self.assertEqual(self.state()["selections"][0]["custom_fields"], [])

    def test_invalid_fields_do_not_change_workbook(self):
        project = self.state()["projects"][0]["Project ID"]
        invalid = [[{"name": "", "value": "x"}], [{"name": "Name", "value": "x"}, {"name": " name ", "value": "y"}],
                   [{"name": "Project ID", "value": "x"}], [{"name": "Long", "value": "x" * 2001}],
                   [{"name": "Broken", "value": {"nested": 1}}], "not a list", [{"name": "X", "value": "\u0000"}]]
        for fields in invalid:
            before = self.path.read_bytes()
            response = self.post("/api/projects", {"values": {"Project ID": project}, "custom_fields": fields})
            self.assertEqual(response.status_code, 400, response.json)
            self.assertEqual(self.path.read_bytes(), before)

    def test_custom_edits_do_not_queue_product_search(self):
        engine = AutoLookup(self.store, lambda: {"ready": True, "name": "brave", "key_name": "BRAVE_API_KEY"}, search=Mock())
        engine.observe(self.store.snapshot()[0])
        with patch.object(app, "automation", engine):
            response = self.post("/api/selections", {"row": self.state()["selections"][0]["_row"], "values": {}, "custom_fields": [{"name": "Supplier", "value": "New supplier"}]})
            self.assertEqual(response.status_code, 200, response.json)
            self.assertFalse(response.json["queued"])
            engine.tick()
            engine.search.assert_not_called()

    def test_custom_fields_in_both_presentations_without_cross_project_leak(self):
        from pptx import Presentation
        try:
            from pypdf import PdfReader
        except ImportError:
            self.skipTest("Install requirements-dev.txt for PDF text verification.")
        data = self.state()
        project, other = data["projects"][:2]
        row = next(r for r in data["selections"] if r["Project ID"] == project["Project ID"])
        long_value = "Material specification " * 70 + "END-OF-CUSTOM-VALUE"
        self.post("/api/projects", {"values": {"Project ID": project["Project ID"]}, "custom_fields": [{"name": "Lot number", "value": "Lot 24"}, {"name": "Long specification", "value": long_value}]})
        self.post("/api/projects", {"values": {"Project ID": other["Project ID"]}, "custom_fields": [{"name": "Private other project", "value": "OTHER-PROJECT-ONLY"}]})
        self.post("/api/selections", {"row": row["_row"], "values": {}, "custom_fields": [{"name": "Supplier", "value": "Approved supplier"}]})
        pdf = self.post("/api/presentation", {"project": project["Project ID"], "mode": "draft", "format": "pdf"})
        self.assertEqual(pdf.status_code, 200, pdf.json if pdf.is_json else "")
        text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.data)).pages)
        for value in ["Lot number", "Lot 24", "Approved supplier", "END-OF-CUSTOM-VALUE"]:
            self.assertIn(value, text)
        self.assertNotIn("OTHER-PROJECT-ONLY", text)
        pptx = self.post("/api/presentation", {"project": project["Project ID"], "mode": "draft", "format": "pptx"})
        self.assertEqual(pptx.status_code, 200, pptx.json if pptx.is_json else "")
        prs = Presentation(io.BytesIO(pptx.data))
        text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
        for value in ["Lot number", "Lot 24", "Approved supplier", "END-OF-CUSTOM-VALUE"]:
            self.assertIn(value, text)
        self.assertNotIn("OTHER-PROJECT-ONLY", text)

    def test_malformed_custom_fields_import_is_rejected(self):
        wb, _ = self.store.snapshot()
        ws = wb["Projects"]
        column = ws.max_column + 1
        ws.cell(1, column, "Custom Fields")
        ws.cell(2, column, "not JSON")
        output = io.BytesIO()
        wb.save(output)
        before = self.path.read_bytes()
        response = self.client.post("/api/workbook/import", data={"file": (io.BytesIO(output.getvalue()), "tracker.xlsx"), "revision": self.state()["revision"], "confirm": "replace-with-backup"}, headers=self.headers)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.path.read_bytes(), before)

    def test_formula_like_values_are_literal(self):
        project = self.state()["projects"][0]["Project ID"]
        fields = [{"name": "Calculation", "value": "=1+1"}, {"name": "Markup", "value": "<script>alert(1)</script>"}]
        response = self.post("/api/projects", {"values": {"Project ID": project}, "custom_fields": fields})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.state()["projects"][0]["custom_fields"], fields)
        wb, _ = self.store.snapshot()
        sheet = common.Sheet(wb["Projects"])
        self.assertEqual(sheet.ws.cell(2, sheet.cols["Custom Fields"]).data_type, "s")


if __name__ == "__main__":
    unittest.main()
