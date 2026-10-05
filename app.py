import datetime as dt
import io
import json
import os
import re
import secrets
import tempfile
import threading
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file
from werkzeug.exceptions import HTTPException

import build_lookbook
import common
import find_urls
from auto_lookup import AutoLookup
from workbook_store import (CLIENT_STATUSES, ConflictError, IDENTITY_FIELDS, OPTIONAL_SELECTION_FIELDS,
                            PROJECT_FIELDS, PROJECT_FORMULAS, SELECTION_FIELDS, STATUSES, ValidationError,
                            WorkbookStore, delete_record, ensure_selection_columns, http_url, next_row,
                            put, records, stale_check, validate_values, detailed_records, set_custom_fields)

BASE = common.BASE_DIR
app = Flask(__name__, template_folder=str(BASE), static_folder=None)
app.config.update(MAX_CONTENT_LENGTH=8 * 1024 * 1024, TRUSTED_HOSTS=["localhost", "127.0.0.1", "[::1]"])
store = WorkbookStore(common.WORKBOOK)
TOKEN = secrets.token_urlsafe(32)
EXPORT_LOCK = threading.Lock()
automation = None


@app.before_request
def guard():
    if request.method not in {"GET", "HEAD", "OPTIONS"}:
        if not secrets.compare_digest(request.headers.get("X-UH-Token", ""), TOKEN):
            return jsonify(error="This page has expired. Refresh and try again."), 403


@app.after_request
def headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Cache-Control"] = "no-store"
    response.headers["Content-Security-Policy"] = "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' https: http: data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'"
    return response


@app.errorhandler(Exception)
def error_response(error):
    if isinstance(error, HTTPException):
        return jsonify(error="The request or file is too large." if error.code == 413 else error.description), error.code
    if isinstance(error, ConflictError):
        return jsonify(error=str(error)), 409
    if isinstance(error, ValidationError):
        return jsonify(error=str(error)), 400
    if isinstance(error, PermissionError):
        return jsonify(error="Close the master workbook in Excel, then retry. Your existing workbook has not been replaced."), 409
    if isinstance(error, FileNotFoundError):
        return jsonify(error="A required file is missing. Check the workbook and template configuration."), 404
    app.logger.error("Request failed: %s", type(error).__name__)
    return jsonify(error="The operation could not be completed. Your source workbook is preserved; check the terminal and retry."), 500


@app.get("/")
def index():
    return render_template("ui.html", token=TOKEN)


@app.get("/favicon.ico")
def favicon():
    return send_file(BASE / "brand_assets" / "uh-icon.png", mimetype="image/png")


@app.get("/brand/<name>.png")
def brand_asset(name):
    sources = {"wordmark": build_lookbook.CFG.get("logo_path"),
               "wordmark-print": build_lookbook.CFG.get("logo_print_path"),
               "symbol": build_lookbook.CFG.get("logo_mark_path"),
               "icon": "brand_assets/uh-icon.png"}
    source = sources.get(name)
    if not source:
        return jsonify(error="Brand asset not found."), 404
    path = Path(source)
    return send_file(path if path.is_absolute() else BASE / path)


@app.get("/assets/<name>")
def asset(name):
    if name not in {"ui.js", "ui.css"}:
        return jsonify(error="Asset not found."), 404
    return send_file(BASE / name)


def provider_info():
    provider = os.getenv("SEARCH_PROVIDER", "brave").lower()
    keys = {"brave": ["BRAVE_API_KEY"], "serpapi": ["SERPAPI_KEY", "SERPAPI_API_KEY"],
            "claude": ["ANTHROPIC_API_KEY"]}.get(provider, [])
    return {"name": provider, "ready": any(os.getenv(key) for key in keys),
            "key_name": keys[0] if keys else "SEARCH_PROVIDER"}


@app.get("/api/state")
def state():
    with store.lock:
        wb, revision = store.snapshot()
        rows = detailed_records(wb, "Selections")
        if automation:
            automation.observe(wb)
            rows = automation.display_rows(rows)
        rows = [{**row, "_stale": stale_check(row)} for row in rows]
        return jsonify(projects=detailed_records(wb, "Projects"), selections=rows,
                       manufacturers=records(wb, "Manufacturers"), sections=common.section_order(wb),
                       statuses=STATUSES, client_statuses=CLIENT_STATUSES, stale_days=180,
                       revision=revision, provider=provider_info(),
                       automation=automation.status() if automation else {"enabled": False, "pending": [], "message": ""},
                       workbook=store.path.name, branding=json.loads((BASE / "lookbook_config.json").read_text()))


@app.post("/api/automation/retry")
def retry_automation():
    current(payload())
    if automation:
        automation.retry()
    return jsonify(ok=True)


def payload():
    data = request.get_json()
    if not isinstance(data, dict) or not isinstance(data.get("revision"), str):
        raise ValidationError("A workbook revision is required. Refresh this page.")
    return data


def current(data):
    wb, revision = store.snapshot()
    if data["revision"] != revision:
        raise ConflictError("The workbook changed. Refresh before saving your changes.")
    if automation:
        automation.observe(wb)
    return wb, revision


@app.post("/api/projects")
def save_project():
    data = payload()
    values = validate_values(data.get("values"), PROJECT_FIELDS)
    with store.lock:
        wb, revision = current(data)
        sheet = common.Sheet(wb["Projects"])
        existing = next((p for p in records(wb, "Projects") if p["Project ID"] == values.get("Project ID")), None)
        if data.get("delete"):
            if not existing:
                raise ConflictError("This project no longer exists. Refresh the page.")
            if any(r["Project ID"] == existing["Project ID"] for r in records(wb, "Selections")):
                raise ValidationError("This project still has selections. Remove them first, then delete the project.")
            delete_record(wb, "Projects", existing["_row"])
            store.save(wb, revision, f"delete project {existing['Project ID']}")
            return jsonify(ok=True, deleted=True)
        if data.get("create") and existing:
            raise ValidationError("That Project ID already exists. Choose a unique ID.")
        row = existing["_row"] if existing else next_row(sheet)
        put(sheet, row, values)
        for field, formula in PROJECT_FORMULAS.items():
            if field in sheet.cols:
                sheet.ws.cell(row, sheet.cols[field], formula.format(row=row))
        if "custom_fields" in data:
            set_custom_fields(wb, "Projects", row, data["custom_fields"])
        store.save(wb, revision, f"project {values.get('Project ID')}")
    return jsonify(ok=True, project_id=values.get("Project ID"))


@app.post("/api/selections")
def save_selection():
    data = payload()
    values = validate_values(data.get("values"), SELECTION_FIELDS + OPTIONAL_SELECTION_FIELDS)
    with store.lock:
        wb, revision = current(data)
        sheet = common.Sheet(wb["Selections"])
        if data.get("delete"):
            old = next((r for r in records(wb, "Selections") if r["_row"] == data.get("row")), None)
            if not old:
                raise ConflictError("This selection no longer exists. Refresh the page.")
            delete_record(wb, "Selections", old["_row"])
            store.save(wb, revision, f"delete selection row {old['_row']} ({old['Item']})")
            if automation:
                automation.observe(wb)
            return jsonify(ok=True, deleted=True)
        old = None
        if data.get("row") is not None:
            old = next((r for r in records(wb, "Selections") if r["_row"] == data["row"]), None)
            if not old:
                raise ConflictError("This selection no longer exists. Refresh the page.")
        row = old["_row"] if old else next_row(sheet)
        merged = {k: old.get(k, "") if old else "" for k in SELECTION_FIELDS + OPTIONAL_SELECTION_FIELDS}
        merged.update(values)
        if not merged["Project ID"] or not merged["Item"]:
            raise ValidationError("Project and item name are required.")
        if automation and str(row) in automation.pending and values.get("Lookup Status") == "Verified":
            raise ValidationError("Automatic lookup is pending for this changed product. Wait for the result or save a new manual link before verifying.")
        merged["Client Status"] = merged.get("Client Status") or "Proposed"
        if old and any(merged[k] != old[k] for k in IDENTITY_FIELDS):
            if merged["Client Status"] in {"Presented", "Approved"}:
                merged["Client Status"] = "Changed"
            elif merged["Client Status"] == "Rejected":
                merged["Client Status"] = "Proposed"
            new_url = merged["Product URL"] if merged["Product URL"] != old["Product URL"] else ""
            merged.update({k: "" for k in ["Product URL", "Image URL", "Checked On", "Lookup Notes"]})
            if merged["Product Name"] == old["Product Name"]:
                merged["Product Name"] = ""
            merged["Product URL"] = new_url
            merged["Lookup Status"] = "Found - verify" if new_url else "Not run"
        elif old and merged["Product URL"] != old["Product URL"]:
            merged["Lookup Status"] = "Found - verify" if merged["Product URL"] else "Not run"
            merged["Checked On"] = ""
            merged["Lookup Notes"] = "Manually updated link. Review before verifying."
        elif not old and (merged["Lookup Status"] == "Verified" or merged["Product URL"]):
            merged["Lookup Status"] = "Found - verify"
        merged["Lookup Status"] = merged["Lookup Status"] or "Not run"
        merged["Include in Lookbook"] = merged["Include in Lookbook"] or "Yes"
        if merged["Lookup Status"] == "Verified" and (not old or old["Lookup Status"] != "Verified"):
            if not data.get("confirm_verified"):
                raise ValidationError("Confirm that you checked the manufacturer, model, and finish on the product page.")
            merged["Checked On"] = dt.date.today().isoformat()
        sheet = ensure_selection_columns(wb)
        put(sheet, row, merged)
        from openpyxl.utils import get_column_letter
        sheet.ws.auto_filter.ref = f"A1:{get_column_letter(sheet.ws.max_column)}{max(sheet.ws.max_row, row)}"
        if "custom_fields" in data:
            set_custom_fields(wb, "Selections", row, data["custom_fields"])
        store.save(wb, revision, f"selection row {row}")
        if automation:
            automation.observe(wb)
        queued = bool(automation and str(row) in automation.pending)
    return jsonify(ok=True, row=row, queued=queued)


@app.post("/api/manufacturers")
def save_manufacturer():
    data = payload()
    values = validate_values(data.get("values"), ["Manufacturer", "Official Domain", "Notes"])
    name = values.get("Manufacturer", "")
    with store.lock:
        wb, revision = current(data)
        sheet = common.Sheet(wb["Manufacturers"])
        old = next((r for r in records(wb, "Manufacturers") if r["Manufacturer"].lower() == name.lower()), None)
        if data.get("delete"):
            if not old:
                raise ConflictError("This manufacturer no longer exists. Refresh the page.")
            delete_record(wb, "Manufacturers", old["_row"])
            store.save(wb, revision, f"delete manufacturer {name}")
            return jsonify(ok=True, deleted=True)
        domain = values.get("Official Domain", "").lower().removeprefix("https://").removeprefix("http://").removeprefix("www.").rstrip("/")
        if not name or not re.fullmatch(r"[a-z0-9](?:[a-z0-9.-]*[a-z0-9])?\.[a-z]{2,}", domain):
            raise ValidationError("Enter a manufacturer name and a domain such as brand.com, without a page path.")
        values["Official Domain"] = domain
        put(sheet, old["_row"] if old else next_row(sheet), values)
        store.save(wb, revision, f"manufacturer {name}")
    return jsonify(ok=True)


@app.post("/api/selections/<int:row>/lookup")
def lookup_selection(row):
    data = payload()
    info = provider_info()
    if not info["ready"]:
        raise ValidationError(f"Product search is not configured. Add {info['key_name']} to your local .env file and restart the app. Do not paste keys into the workbook or chat.")
    wb, revision = current(data)
    record = next((r for r in records(wb, "Selections") if r["_row"] == row), None)
    if not record:
        raise ValidationError("Selection not found.")
    if automation and str(row) in automation.pending:
        raise ConflictError("Automatic lookup is already queued for this product. Check its progress or use Retry automatic lookup.")
    if record["Lookup Status"] == "Verified":
        return jsonify(ok=True, skipped=True)
    if not record["Manufacturer"] or not record["Model #"]:
        raise ValidationError("Add both a manufacturer and model number before searching.")
    domain = common.manufacturer_domains(wb).get(record["Manufacturer"].lower(), "")
    result = find_urls.lookup(find_urls.PROVIDERS[info["name"]], record["Manufacturer"], record["Model #"], domain,
                              product_name=record["Product Name"] or record["Item"])
    if result["status"] == "Search error":
        return jsonify(error=result["notes"]), 502
    if len(result.get("url", "")) > 4000:
        raise ValidationError("The returned product URL is too long. Review this product manually.")
    updates = {"Lookup Status": result["status"], "Lookup Notes": result.get("notes", ""), "Checked On": dt.date.today().isoformat()}
    if result.get("url"):
        updates["Product URL"] = result["url"]
        updates["Product Name"] = result.get("name", "")
        if not record["Image URL"] or record["Image URL"].startswith("http"):
            updates["Image URL"] = result.get("image", "")
    elif record["Product URL"]:
        updates["Lookup Notes"] += " Existing link retained for manual review; it was not confirmed."
    with store.lock:
        put(common.Sheet(wb["Selections"]), row, updates)
        store.save(wb, revision, f"lookup row {row}")
        if automation:
            automation.acknowledge(wb, row)
    return jsonify(ok=True, result=result)


@app.post("/api/selections/parse-link")
def parse_link():
    data = payload()
    current(data)
    url = common.clean(data.get("url"))
    if not http_url(url) or len(url) > 4000:
        raise ValidationError("Enter a full HTTP or HTTPS product link.")
    details = find_urls.describe_page(url)
    if not details or not details.get("Product Name"):
        raise ValidationError("That page could not be read as a product page. Enter the details manually.")
    return jsonify(ok=True, details=details)


@app.get("/api/workbook")
def export_workbook():
    with store.lock:
        raw = store.path.read_bytes()
    return send_file(io.BytesIO(raw), as_attachment=True, download_name=store.path.name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/workbook/import")
def import_workbook():
    if request.form.get("confirm") != "replace-with-backup":
        raise ValidationError("Confirm replacement of the master workbook first.")
    upload = request.files.get("file")
    if not upload or not upload.filename.lower().endswith(".xlsx"):
        raise ValidationError("Choose an .xlsx workbook exported from this app or using the same template.")
    with store.lock:
        store.import_bytes(upload.read(), request.form.get("revision", ""))
        if automation:
            automation.observe(store.snapshot()[0])
    return jsonify(ok=True)


def schedule_options(data):
    """Per-export schedule title and item code prefix, falling back to the branding defaults."""
    title, prefix = common.clean(data.get("title")), common.clean(data.get("prefix"))
    if len(title) > 80 or any(ord(character) < 32 for character in title):
        raise ValidationError("Enter a schedule title of up to 80 characters on one line.")
    if prefix and not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9-]{0,5}", prefix):
        raise ValidationError("The item code prefix must be 1-6 letters, numbers, or hyphens, such as EX.")
    return title or None, prefix or None


@app.post("/api/presentation")
def presentation():
    data = payload()
    wb, revision = current(data)
    project = next((p for p in detailed_records(wb, "Projects") if p["Project ID"] == data.get("project")), None)
    if not project:
        raise ValidationError("Choose a valid project.")
    mode, kind = data.get("mode", "draft"), data.get("format", "pdf")
    if mode not in {"draft", "verified", "final"} or kind not in {"pdf", "pptx"}:
        raise ValidationError("Invalid presentation options.")
    title, prefix = schedule_options(data)
    rows = [r for r in detailed_records(wb, "Selections") if r["Project ID"] == project["Project ID"] and r["Include in Lookbook"].lower() != "no"]
    if automation:
        with store.lock:
            pending_rows = set(automation.pending)
        if mode == "verified":
            rows = [r for r in rows if str(r["_row"]) not in pending_rows]
        elif any(str(r["_row"]) in pending_rows for r in rows):
            raise ConflictError("This project has changed products awaiting automatic lookup. Close Excel and wait for the results, save manual links, or export verified selections only.")
    if mode == "final" and any(r["Lookup Status"] != "Verified" for r in rows):
        raise ValidationError("Some included selections are not verified. Review them, export a draft, or choose verified selections only.")
    if mode == "verified":
        rows = [r for r in rows if r["Lookup Status"] == "Verified"]
    if not rows:
        raise ValidationError("No selections match these export options.")
    order = common.section_order(wb)
    rows.sort(key=lambda r: (order.index(r["Section"]) if r["Section"] in order else len(order), r["Section"], r["_row"]))
    note = f"export {project['Project ID']} {kind} {mode}"
    with store.lock:
        marked = []
        if mode != "draft":
            sheet = ensure_selection_columns(wb)
            marked = [r for r in rows if common.clean(r.get("Client Status")) in {"", "Proposed"}]
            for r in marked:
                put(sheet, r["_row"], {"Client Status": "Presented"})
        if marked:
            store.save(wb, revision, note)
        else:
            store.log(note)
        rev = store.count_notes(f"export {project['Project ID']} ")
    with EXPORT_LOCK, tempfile.TemporaryDirectory(prefix="uh_presentation_") as folder:
        if kind == "pdf":
            build_schedule, mime = build_lookbook.build, "application/pdf"
        else:
            from build_powerpoint import build as build_schedule
            mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        output = build_schedule(project["Project ID"], project, rows, mode == "draft",
                                output_dir=Path(folder), title=title, prefix=prefix, rev=rev)
        raw = output.read_bytes()
    return send_file(io.BytesIO(raw), as_attachment=True, download_name=output.name, mimetype=mime)


if __name__ == "__main__":
    import argparse
    import atexit
    parser = argparse.ArgumentParser(description="Run the local UH Homes selection tracker.")
    parser.add_argument("--port", type=int, default=5000)
    args = parser.parse_args()
    if not 1 <= args.port <= 65535:
        parser.error("Port must be between 1 and 65535.")
    automation = AutoLookup(store, provider_info)
    automation.observe(store.snapshot()[0])
    automation.start()
    atexit.register(automation.stop)
    print("Automatic lookup watches saved app and Excel edits while this process is running.")
    print(f"UH Homes Selections: http://127.0.0.1:{args.port}")
    print("Local access only. Close the master workbook in Excel before saving changes.")
    app.run(host="127.0.0.1", port=args.port, debug=False)
