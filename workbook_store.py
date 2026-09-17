import datetime as dt
import hashlib
import io
import os
import re
import shutil
import tempfile
import threading
import zipfile
from copy import copy
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

from common import Sheet, clean

PROJECT_FIELDS = ["Project ID", "Project Name", "Client Name", "Address", "Plan / Elevation", "Designer", "Presentation Date", "Cover Image"]
SELECTION_FIELDS = ["Project ID", "Section", "Room / Area", "Item", "Manufacturer", "Model #", "Finish / Color", "Qty", "Product URL", "Product Name", "Image URL", "Lookup Status", "Checked On", "Lookup Notes", "Client Notes", "Include in Lookbook"]
STATUSES = ["Not run", "Found - verify", "Found - retailer", "Multiple matches", "Not found", "Search error", "Verified"]
IDENTITY_FIELDS = ["Item", "Product Name", "Manufacturer", "Model #", "Finish / Color"]


class ValidationError(ValueError):
    pass


class ConflictError(ValueError):
    pass


def http_url(value):
    try:
        parsed = urlsplit(value)
        return parsed.scheme in {"http", "https"} and bool(parsed.hostname) and not parsed.username and not parsed.password
    except ValueError:
        return False


def validate_values(data, fields):
    if not isinstance(data, dict) or set(data) - set(fields):
        raise ValidationError("Unexpected fields in the submitted data.")
    result = {}
    for key, value in data.items():
        if not isinstance(value, (str, int, float, type(None))):
            raise ValidationError(f"Invalid value for {key}.")
        value = clean(value)
        if len(value) > 4000 or any(ord(c) < 32 and c not in "\n\r\t" for c in value):
            raise ValidationError(f"{key} contains invalid or overly long text.")
        if key == "Product URL" and value and not http_url(value):
            raise ValidationError("Product URL must be a full HTTP or HTTPS link.")
        if key == "Qty" and value:
            try:
                qty = float(value)
                if not 0 <= qty <= 1_000_000:
                    raise ValueError()
            except ValueError:
                raise ValidationError("Quantity must be a number between 0 and 1,000,000.") from None
        result[key] = value
    return result


def put(sheet, row, data):
    for key, value in data.items():
        cell = sheet.ws.cell(row, sheet.cols[key])
        cell.value = value or None
        if isinstance(value, str) and value:
            cell.data_type = "s"
        if key == "Product URL":
            cell.hyperlink = value or None


def records(wb, name):
    return [{**{k: v.isoformat() if isinstance(v, (dt.datetime, dt.date)) else clean(v) for k, v in data.items()}, "_row": row}
            for row, data in Sheet(wb[name]).rows()]


def next_row(sheet):
    used = {row for row, _ in sheet.rows()}
    row = 2
    while row in used:
        row += 1
    if row > 10001:
        raise ValidationError("This local workbook supports up to 10,000 rows per sheet.")
    if row > 2:
        for col in sheet.cols.values():
            sheet.ws.cell(row, col)._style = copy(sheet.ws.cell(2, col)._style)
    return row


def validate_workbook(wb):
    required = {"Projects": PROJECT_FIELDS, "Selections": SELECTION_FIELDS,
                "Manufacturers": ["Manufacturer", "Official Domain"], "Lists": ["Sections (lookbook order)"]}
    for name, fields in required.items():
        if name not in wb:
            raise ValidationError(f"Missing worksheet: {name}.")
        ws = wb[name]
        if ws.max_row > 10001 or ws.max_column > 50:
            raise ValidationError(f"Worksheet {name} is too large (maximum 10,000 rows and 50 columns).")
        headers = [clean(c.value) for c in ws[1] if clean(c.value)]
        if len(set(headers)) != len(headers):
            raise ValidationError(f"Duplicate column names in {name}.")
        missing = set(fields) - set(headers)
        if missing:
            raise ValidationError(f"Missing columns in {name}: {', '.join(sorted(missing))}.")
        sheet = Sheet(ws)
        for row, data in sheet.rows():
            for field in fields:
                if ws.cell(row, sheet.cols[field]).data_type == "f":
                    raise ValidationError(f"Formulas are not allowed in input cells: {name}, row {row}, {field}.")
    projects = records(wb, "Projects")
    ids = [p["Project ID"] for p in projects]
    if any(not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,49}", pid) for pid in ids):
        raise ValidationError("Project IDs must be 1–50 letters, numbers, hyphens, or underscores.")
    if len(ids) != len(set(ids)):
        raise ValidationError("Project IDs must be unique.")
    if any(not p["Project Name"] for p in projects):
        raise ValidationError("Every project needs a name.")
    for row in records(wb, "Selections"):
        if row["Project ID"] not in ids:
            raise ValidationError(f"Selections row {row['_row']} references an unknown Project ID.")
        validate_values({key: row[key] for key in SELECTION_FIELDS}, SELECTION_FIELDS)
        if row["Lookup Status"] and row["Lookup Status"] not in STATUSES:
            raise ValidationError(f"Unknown lookup status in row {row['_row']}.")
        if row["Lookup Status"] == "Verified" and not http_url(row["Product URL"]):
            raise ValidationError(f"Verified row {row['_row']} needs a product URL.")
        if row["Include in Lookbook"].lower() not in {"", "yes", "no"}:
            raise ValidationError(f"Include in Lookbook must be Yes or No in row {row['_row']}.")


class WorkbookStore:
    def __init__(self, path):
        self.path = Path(path)
        self.lock = threading.RLock()

    def snapshot(self):
        with self.lock:
            raw = self.path.read_bytes()
            wb = load_workbook(io.BytesIO(raw), keep_links=False)
            validate_workbook(wb)
            return wb, hashlib.sha256(raw).hexdigest()

    def save(self, wb, revision):
        with self.lock:
            validate_workbook(wb)
            if hashlib.sha256(self.path.read_bytes()).hexdigest() != revision:
                raise ConflictError("The workbook changed. Refresh the page before saving again.")
            backup_dir = self.path.parent / "backups"
            backup_dir.mkdir(exist_ok=True)
            stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            shutil.copy2(self.path, backup_dir / f"{self.path.stem}_{stamp}.xlsx")
            handle, temp = tempfile.mkstemp(suffix=".xlsx", dir=self.path.parent)
            os.close(handle)
            try:
                wb.save(temp)
                if hashlib.sha256(self.path.read_bytes()).hexdigest() != revision:
                    raise ConflictError("The workbook changed during the save. Refresh and retry.")
                os.replace(temp, self.path)
            finally:
                Path(temp).unlink(missing_ok=True)

    def import_bytes(self, raw, revision):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 50_000_000:
                    raise ValidationError("The uncompressed workbook exceeds 50 MB.")
                if any(info.filename.lower().endswith("vbaproject.bin") for info in archive.infolist()):
                    raise ValidationError("Macro-enabled workbooks are not supported.")
            wb = load_workbook(io.BytesIO(raw), keep_links=False)
        except ValidationError:
            raise
        except Exception:
            raise ValidationError("This file could not be read as an Excel .xlsx workbook.") from None
        self.save(wb, revision)
