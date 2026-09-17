"""Creates the UH Homes Selections Tracker workbook (run once)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

F = "Arial"
HDR_FILL = PatternFill("solid", start_color="2F3437")
INPUT_FILL = PatternFill("solid", start_color="FFF6D5")   # yellow-ish = you type here
AUTO_FILL = PatternFill("solid", start_color="EEF1F3")    # grey = script fills
thin = Side(style="thin", color="D0D4D8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

SECTIONS = ["Exterior", "Roofing", "Windows & Doors", "Kitchen", "Appliances",
            "Bathrooms", "Flooring", "Lighting", "Paint", "Interior Doors & Trim",
            "Hardware", "Other"]
STATUSES = ["Not run", "Found - verify", "Found - retailer", "Multiple matches",
            "Not found", "Verified"]

wb = Workbook()

# ---------- Instructions ----------
ins = wb.active
ins.title = "Instructions"
lines = [
    ("UH Homes - Selections Tracker", True, 16),
    ("", False, 10),
    ("How this workbook works", True, 12),
    ("1. Add each project once on the Projects sheet (one row per project, unique Project ID).", False, 10),
    ("2. Add every selection on the Selections sheet. Yellow columns are yours to fill; grey columns are filled by find_urls.py.", False, 10),
    ("3. Run:  python find_urls.py   - fills Product URL, Product Name, Image URL, Lookup Status, Checked On, Lookup Notes.", False, 10),
    ("4. Review every row. When a link is confirmed correct, set Lookup Status to 'Verified'. Verified rows are never overwritten.", False, 10),
    ("5. Run:  python build_lookbook.py --project UH-101   (or --all) - creates the client PDF lookbook in /output.", False, 10),
    ("", False, 10),
    ("Legend", True, 12),
    ("Yellow cell = input you type", False, 10),
    ("Grey cell = filled by the scripts (you may edit by hand; set status to Verified to lock it)", False, 10),
    ("", False, 10),
    ("Rules", True, 12),
    ("- Section values come from the Lists sheet; lookbook sections appear in the order listed there.", False, 10),
    ("- Add each brand's official website on the Manufacturers sheet - the lookup searches that site first.", False, 10),
    ("- Image URL may be a web link or a local file path (e.g. images/front-door.jpg) to override the found image.", False, 10),
    ("- Include in Lookbook = No hides a row from the client PDF without deleting it.", False, 10),
    ("- The rows on Projects/Selections are EXAMPLES - replace them with real data.", False, 10),
]
for i, (t, b, s) in enumerate(lines, start=1):
    c = ins.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=b, size=s)
ins["A11"].fill = INPUT_FILL
ins["A12"].fill = AUTO_FILL
ins.column_dimensions["A"].width = 120

def header(ws, cols, widths):
    for i, (h, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

def style_block(ws, rows, ncols, input_cols, auto_cols, formula_cols=()):
    for r in range(2, rows + 2):
        for col in range(1, ncols + 1):
            c = ws.cell(row=r, column=col)
            c.font = Font(name=F, size=10)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if col in input_cols:
                c.fill = INPUT_FILL
            elif col in auto_cols:
                c.fill = AUTO_FILL

# ---------- Projects ----------
pj = wb.create_sheet("Projects")
pcols = ["Project ID", "Project Name", "Client Name", "Address", "Plan / Elevation",
         "Designer", "Presentation Date", "Total Items", "Links Found", "Verified", "Needs Review", "Cover Image"]
header(pj, pcols, [12, 26, 22, 34, 18, 18, 16, 11, 11, 10, 12, 30])
projects = [
    ["UH-101", "The Magnolia Residence", "Example Client A", "123 Example Ln, Frisco, TX", "Plan 2840 / Elev. B", "Designer Name", "2026-10-01"],
    ["UH-102", "Cedar Ridge Custom Build", "Example Client B", "456 Sample Dr, Prosper, TX", "Plan 3310 / Elev. A", "Designer Name", "2026-10-15"],
]
NP = 50
for r, row in enumerate(projects, start=2):
    for c, v in enumerate(row, start=1):
        pj.cell(row=r, column=c, value=v)
for r in range(2, NP + 2):
    pj[f"H{r}"] = f'=IF(A{r}="","",COUNTIF(Selections!$A:$A,A{r}))'
    pj[f"I{r}"] = f'=IF(A{r}="","",COUNTIFS(Selections!$A:$A,A{r},Selections!$I:$I,"?*"))'
    pj[f"J{r}"] = f'=IF(A{r}="","",COUNTIFS(Selections!$A:$A,A{r},Selections!$L:$L,"Verified"))'
    pj[f"K{r}"] = f'=IF(A{r}="","",H{r}-J{r})'
style_block(pj, NP, len(pcols), input_cols=[1, 2, 3, 4, 5, 6, 7, 12], auto_cols=range(8, 12))
pj["L1"].comment = Comment("Optional: web link or local path to a rendering/photo for the lookbook cover.", "Tracker")
pj["H1"].comment = Comment("Formulas - count rows on Selections for this Project ID.", "Tracker")

# ---------- Selections ----------
sl = wb.create_sheet("Selections")
scols = ["Project ID", "Section", "Room / Area", "Item", "Manufacturer", "Model #",
         "Finish / Color", "Qty", "Product URL", "Product Name", "Image URL",
         "Lookup Status", "Checked On", "Lookup Notes", "Client Notes", "Include in Lookbook"]
header(sl, scols, [11, 18, 16, 22, 18, 18, 18, 6, 40, 30, 30, 16, 12, 30, 26, 11])
examples = [
    ["UH-101", "Exterior", "Front Elevation", "Siding", "James Hardie", "HardiePlank Lap", "Arctic White", "", "", "", "", "Not run", "", "", "", "Yes"],
    ["UH-101", "Windows & Doors", "Front Entry", "Entry Door", "Therma-Tru", "S4108", "Black", 1, "", "", "", "Not run", "", "", "", "Yes"],
    ["UH-101", "Kitchen", "Kitchen", "Kitchen Faucet", "Moen", "7594ESRS", "Spot Resist Stainless", 1, "", "", "", "Not run", "", "", "", "Yes"],
    ["UH-101", "Bathrooms", "Primary Bath", "Vanity Faucet", "Delta", "559LF-PP", "Chrome", 2, "", "", "", "Not run", "", "", "", "Yes"],
    ["UH-101", "Paint", "Whole House", "Wall Paint", "Sherwin-Williams", "SW 7005", "Pure White", "", "", "", "", "Not run", "", "", "", "Yes"],
    ["UH-102", "Kitchen", "Kitchen", "Kitchen Sink", "Kohler", "K-5540", "White", 1, "", "", "", "Not run", "", "", "", "Yes"],
    ["UH-102", "Lighting", "Dining Room", "Chandelier", "Kichler", "", "Olde Bronze", 1, "", "", "", "Not run", "", "", "Model TBD", "Yes"],
]
NS = 500
for r, row in enumerate(examples, start=2):
    for c, v in enumerate(row, start=1):
        sl.cell(row=r, column=c, value=v)
style_block(sl, NS, len(scols), input_cols=[1, 2, 3, 4, 5, 6, 7, 8, 15, 16], auto_cols=[9, 10, 11, 12, 13, 14])
sl["F1"].comment = Comment("Type the model / SKU exactly as the manufacturer lists it.", "Tracker")
sl["L1"].comment = Comment("Set to 'Verified' after you confirm the link. Verified rows are never overwritten by find_urls.py.", "Tracker")
sl.auto_filter.ref = f"A1:P{NS + 1}"

# ---------- Manufacturers ----------
mf = wb.create_sheet("Manufacturers")
mcols = ["Manufacturer", "Official Domain", "Notes"]
header(mf, mcols, [24, 30, 50])
brands = [
    ("James Hardie", "jameshardie.com"), ("Therma-Tru", "thermatru.com"), ("Pella", "pella.com"),
    ("Andersen", "andersenwindows.com"), ("Moen", "moen.com"), ("Delta", "deltafaucet.com"),
    ("Kohler", "kohler.com"), ("Sherwin-Williams", "sherwin-williams.com"), ("Behr", "behr.com"),
    ("Kichler", "kichler.com"), ("Daltile", "daltile.com"), ("Shaw", "shawfloors.com"),
    ("GE Appliances", "geappliances.com"), ("Whirlpool", "whirlpool.com"),
    ("GAF", "gaf.com"), ("Owens Corning", "owenscorning.com"), ("Kwikset", "kwikset.com"),
    ("Schlage", "schlage.com"), ("Emtek", "emtek.com"), ("Trex", "trex.com"),
]
for r, (n, d) in enumerate(brands, start=2):
    mf.cell(row=r, column=1, value=n)
    mf.cell(row=r, column=2, value=d)
mf.cell(row=2, column=3, value="Starter list - double-check each domain and add the brands UH Homes uses.")
style_block(mf, 100, 3, input_cols=[1, 2, 3], auto_cols=[])

# ---------- Lists ----------
ls = wb.create_sheet("Lists")
header(ls, ["Sections (lookbook order)", "Lookup Status", "Yes / No"], [26, 20, 10])
for i, s in enumerate(SECTIONS, start=2):
    ls.cell(row=i, column=1, value=s)
for i, s in enumerate(STATUSES, start=2):
    ls.cell(row=i, column=2, value=s)
ls["C2"], ls["C3"] = "Yes", "No"
style_block(ls, 30, 3, input_cols=[1], auto_cols=[])

# ---------- Validations ----------
def dv(ws, formula, rng):
    v = DataValidation(type="list", formula1=formula, allow_blank=True)
    v.error = "Pick a value from the list (edit the Lists sheet to add options)."
    ws.add_data_validation(v)
    v.add(rng)
dv(sl, "=Lists!$A$2:$A$31", f"B2:B{NS + 1}")
dv(sl, "=Lists!$B$2:$B$7", f"L2:L{NS + 1}")
dv(sl, "=Lists!$C$2:$C$3", f"P2:P{NS + 1}")
dv(sl, f"=Projects!$A$2:$A${NP + 1}", f"A2:A{NS + 1}")


wb.save("UH_Homes_Selections_Tracker.xlsx")
print("saved")
