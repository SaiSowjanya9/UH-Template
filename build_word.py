"""
Editable Word twin of the client finish schedule.

Usage:
    python build_word.py --project UH-101
    python build_word.py --project UH-101 --draft
    python build_word.py --project UH-101 --fillable
    python build_word.py --project UH-101 --spec
    python build_word.py --project UH-101 --title "Interior Selections" --prefix IN
    python build_word.py --project UH-101 --hide "code,Qty,Product URL"

Output: output/<ProjectID>_<Project Name>_<Schedule title>.docx

The page keeps the approved schedule's geometry - letter portrait, the same margins, type
sizes and column widths, the wordmark header, client band, letter-spaced section headings,
coded rows and sign-off - but every block is a real Word paragraph or table, so Word
repeats the furniture, reflows edited text and repaginates on its own. Section rules are
tab leaders rather than drawn lines for the same reason. Page breaks may therefore differ
from the PDF.

The fillable copy holds the client details, each line's description and a remove tick in
content controls tagged with the workbook row, so a marked-up document imports back through
the same route as the fillable PDF. The spec sheet is the internal landscape column dump and
is the only Word document here that carries pricing.
"""
import argparse
import datetime as dt
import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor
from docx.text.run import Run
from reportlab.pdfbase import pdfmetrics

from build_form import HIDDEN_FIELD, REMOVE_W, form_band_values
from build_lookbook import (BAND_GAP, BAND_H, BODY_BOTTOM, BODY_TOP, CFG, CODE_STYLE, CODE_W,
                            CONTENT_W, DESC_STYLE, FONT_B, FONT_BB, HEADING_OVERHANG, HIDEABLE,
                            ITEM_STYLE, ITEM_W, LOGO_H, M, MARK_W, NOTE_STYLE, OUT, PAD_X, PAD_Y,
                            PAGE_H, PAGE_W, PRODUCT_LINK, ROOM_STYLE, SCHEDULE, SECTION_GAP,
                            SIGN_H, SIGN_W, SIGN_X, STATUS_STYLE, STATUS_W, band_values, clean,
                            code_prefix, description, fit, group_sections, hidden_columns,
                            load_data, numbered, presentation_details, product_link, row_widths,
                            schedule_title, spaced_caps)
from branding import load_logo

COLORS = {key: value.lstrip("#").upper() for key, value in CFG["colors"].items()}
FONTS = CFG.get("fonts", {})
FAMILY = FONTS.get("family") or Path(FONTS.get("body") or "Arial").stem.split("-")[0] or "Arial"
# (font size, leading) taken from the PDF's styles so one change moves both deliverables.
CODE, ITEM, ROOM, DESC, STATUS, NOTE = ((style.fontSize, style.leading) for style in
                                        (CODE_STYLE, ITEM_STYLE, ROOM_STYLE, DESC_STYLE,
                                         STATUS_STYLE, NOTE_STYLE))
HEADER_TOP, FOOTER_BOTTOM = 32, 24   # header and footer insets from the page edges
BAND_INSET = 10.8                    # left and right padding inside a band field
SPEC_TITLE = "Spec Sheet"
# Internal column dump; Section becomes the heading and the lookbook flag a note under the item.
# (column, width share, the hideable key that drops it); shares are rescaled to what is shown.
SPEC_COLUMNS = [("Room / Area", 68, "Room / Area"), ("Item", 100, "Item"),
                ("Manufacturer", 78, "Manufacturer"), ("Model #", 66, "Model #"),
                ("Finish / Color", 80, "Finish / Color"), ("Qty", 28, "Qty"),
                ("Unit Price", 46, "Unit Price"), ("Markup %", 42, "Markup %"),
                ("Client Status", 54, "Client Status"), (PRODUCT_LINK, 52, "Product URL"),
                ("Client Notes", 84, "Client Notes")]

# WordprocessingML fixes the order of these property children, so new ones are inserted in place.
PPR_ORDER = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
             "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs",
             "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr")
TCPR_ORDER = ("w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
              "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
TBLPR_ORDER = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
               "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
               "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
TRPR_ORDER = ("w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore", "w:wAfter",
              "w:cantSplit", "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden")

PAGE_FIELD = ('<w:fldSimple {ns} w:instr=" PAGE "><w:r><w:rPr>'
              '<w:rFonts w:ascii="{family}" w:hAnsi="{family}"/><w:sz w:val="{half}"/>'
              '<w:color w:val="{color}"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
RUN_PROPERTIES = ('<w:rFonts w:ascii="{family}" w:hAnsi="{family}"/><w:sz w:val="{half}"/>'
                  '<w:color w:val="{color}"/>')
TEXT_CONTROL = ('<w:sdt {ns}><w:sdtPr><w:alias w:val="{alias}"/><w:tag w:val="{tag}"/>'
                '<w:id w:val="{ident}"/><w:rPr>{rpr}</w:rPr><w:text/></w:sdtPr><w:sdtContent>'
                '<w:r><w:rPr>{rpr}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'
                '</w:sdtContent></w:sdt>')
CHECK_CONTROL = ('<w:sdt {ns}><w:sdtPr><w:alias w:val="Remove"/><w:tag w:val="{tag}"/>'
                 '<w:id w:val="{ident}"/><w14:checkbox><w14:checked w14:val="0"/>'
                 '<w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
                 '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/></w14:checkbox>'
                 '</w:sdtPr><w:sdtContent><w:r><w:rPr><w:rFonts w:ascii="MS Gothic" '
                 'w:eastAsia="MS Gothic" w:hAnsi="MS Gothic"/><w:sz w:val="20"/></w:rPr>'
                 '<w:t>\u2610</w:t></w:r></w:sdtContent></w:sdt>')
# The diagonal DRAFT mark is a header watermark, the only way Word repeats art behind the text.
WATERMARK = (
    '<w:p {ns} xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
    '<w:pPr><w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/></w:pPr>'
    '<w:r><w:pict><v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800"'
    ' path="m@7,0l@8,0m@5,21600l@6,21600e"><v:formulas>'
    '<v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/>'
    '<v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/>'
    '<v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/><v:f eqn="if @0 @4 21600"/>'
    '<v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/>'
    '<v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas>'
    '<v:path textpathok="t" o:connecttype="custom"'
    ' o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>'
    '<v:textpath on="t" fitshape="t"/></v:shapetype>'
    '<v:shape id="UHDraftWatermark" type="#_x0000_t136" style="position:absolute;margin-left:0;'
    'margin-top:0;width:420pt;height:157pt;rotation:332;z-index:-251656192;'
    'mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
    'mso-position-vertical:center;mso-position-vertical-relative:margin"'
    ' fillcolor="#{accent}" stroked="f"><v:fill opacity=".07"/>'
    '<v:textpath style="font-family:&quot;{family}&quot;;font-size:1pt" string="DRAFT"/>'
    '</v:shape></w:pict></w:r></w:p>')


# --------------------------------------------------------------------------
# Low-level XML helpers
# --------------------------------------------------------------------------
def property_element(parent, tag, order):
    """Fetch or create a property child, keeping the element order the schema demands."""
    node = parent.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        parent.insert_element_before(node, *order[order.index(tag) + 1:])
    return node


def twips(points):
    return str(int(round(points * 20)))


def margins(parent, tag, order, **sides):
    """Cell or table padding, in points, written in the schema's top/left/bottom/right order."""
    node = property_element(parent, tag, order)
    for side, value in sides.items():
        child = OxmlElement(f"w:{side}")
        child.set(qn("w:w"), twips(value))
        child.set(qn("w:type"), "dxa")
        node.append(child)
    return node


def border(edge, width, color):
    node = OxmlElement(f"w:{edge}")
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(max(2, round(width * 8))))   # Word measures rules in eighths of a point
    node.set(qn("w:space"), "0")
    node.set(qn("w:color"), color)
    return node


def rule(paragraph, edge, width, color):
    property_element(paragraph._p.get_or_add_pPr(), "w:pBdr", PPR_ORDER).append(border(edge, width, color))


def cell_rule(cell, edge, width, color):
    property_element(cell._tc.get_or_add_tcPr(), "w:tcBorders", TCPR_ORDER).append(border(edge, width, color))


def shade(cell, color):
    node = property_element(cell._tc.get_or_add_tcPr(), "w:shd", TCPR_ORDER)
    node.set(qn("w:val"), "clear")
    node.set(qn("w:color"), "auto")
    node.set(qn("w:fill"), color)


def row_height(row, points, rule_="atLeast"):
    node = property_element(row._tr.get_or_add_trPr(), "w:trHeight", TRPR_ORDER)
    node.set(qn("w:val"), twips(points))
    node.set(qn("w:hRule"), rule_)


def keep_whole(row, header=False):
    """Stop Word splitting a row across pages, and optionally repeat it as a header."""
    property_element(row._tr.get_or_add_trPr(), "w:cantSplit", TRPR_ORDER)
    if header:
        property_element(row._tr.get_or_add_trPr(), "w:tblHeader", TRPR_ORDER)


# --------------------------------------------------------------------------
# Text and layout helpers
# --------------------------------------------------------------------------
def shape(paragraph, size=None, leading=None, align=None, before=0, after=0, indent=0, keep=False):
    """Exact leading and no stray Word spacing, so the page matches the printed schedule."""
    size = size or DESC[0]
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing = Pt(leading or size * 1.3)
    fmt.left_indent = Pt(indent)
    if align is not None:
        paragraph.alignment = align
    paragraph.keep_with_next = keep
    return paragraph


def run_text(paragraph, text, size, color, bold=False):
    run = paragraph.add_run(str(text or ""))
    run.font.name, run.font.size, run.font.bold = FAMILY, Pt(size), bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def write(paragraph, text, size=None, color=None, bold=False, **geometry):
    size = size or DESC[0]
    shape(paragraph, size=size, **geometry)
    run_text(paragraph, clean(text), size, color or COLORS["body"], bold)
    return paragraph


def picture(paragraph, source, height=None, width=None):
    """Logo at its supplied aspect ratio; returns the drawn width, or 0 when missing."""
    image = load_logo(source)
    if image is None:
        return 0
    stream = io.BytesIO()
    image.save(stream, "PNG")
    stream.seek(0)
    iw, ih = image.size
    height = height if height is not None else width * ih / iw
    width = width if width is not None else height * iw / ih
    paragraph.add_run().add_picture(stream, width=Pt(width), height=Pt(height))
    return width


def run_properties(size, color, vanish=False):
    return RUN_PROPERTIES.format(family=escape(FAMILY, {'"': "&quot;"}), half=int(round(size * 2)),
                                 color=color) + ("<w:vanish/>" if vanish else "")


_identifiers = iter(range(900001, 999999))


def control(paragraph, tag, value, alias="Schedule field", size=None, color=None, vanish=False):
    """A plain-text content control: click once in Word and type, and it reads back by tag."""
    size = size or DESC[0]
    paragraph._p.append(parse_xml(TEXT_CONTROL.format(
        ns=nsdecls("w"), alias=escape(alias, {'"': "&quot;"}), tag=escape(tag, {'"': "&quot;"}),
        ident=next(_identifiers), rpr=run_properties(size, color or COLORS["body"], vanish),
        text=escape(clean(value)))))
    return paragraph


def hyperlink(paragraph, url, text, size, color):
    """A real Word hyperlink, so the printed 'Product Link' opens the product page."""
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    element = OxmlElement("w:r")
    link.append(element)
    paragraph._p.append(link)
    run = Run(element, paragraph)
    run.text = text
    run.font.name, run.font.size, run.font.underline = FAMILY, Pt(size), True
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def checkbox(paragraph, tag):
    paragraph._p.append(parse_xml(CHECK_CONTROL.format(ns=nsdecls("w", "w14"),
                                                       tag=escape(tag, {'"': "&quot;"}),
                                                       ident=next(_identifiers))))
    return paragraph


def grid(document, widths, pad_x=PAD_X, pad_y=PAD_Y, indent=0):
    """A borderless table measured in points, like the PDF's tables."""
    table = document.add_table(rows=0, cols=len(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    property_element(properties, "w:tblLayout", TBLPR_ORDER).set(qn("w:type"), "fixed")
    left = property_element(properties, "w:tblInd", TBLPR_ORDER)
    left.set(qn("w:w"), twips(indent))
    left.set(qn("w:type"), "dxa")
    margins(properties, "w:tblCellMar", TBLPR_ORDER, top=pad_y, left=pad_x, bottom=pad_y, right=pad_x)
    for column, width in zip(table.columns, widths):
        column.width = Pt(width)
    return table


def add_row(table, widths, align=WD_ALIGN_VERTICAL.CENTER):
    cells = table.add_row().cells
    for cell, width in zip(cells, widths):
        cell.width = Pt(width)
        cell.vertical_alignment = align
    return cells


# --------------------------------------------------------------------------
# Page furniture
# --------------------------------------------------------------------------
def plain(paragraph):
    """Word's Header and Footer styles carry their own tab stops, which would fight ours."""
    paragraph.style = "Normal"
    return paragraph


def furniture(section, title, draft, content_w):
    """Header and footer, which Word then repeats on every page."""
    head = section.header
    head.is_linked_to_previous = False
    eyebrow = clean(CFG.get("tagline", "Finish Schedule")).upper()
    write(plain(head.paragraphs[0]), f"{eyebrow}   \u00b7   DRAFT" if draft else eyebrow, 7,
          COLORS["accent"], bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, leading=9)
    banner = shape(plain(head.add_paragraph()), leading=LOGO_H + 2)
    banner.paragraph_format.tab_stops.add_tab_stop(Pt(content_w), WD_TAB_ALIGNMENT.RIGHT)
    rule(banner, "bottom", 0.8, COLORS["dark"])
    drawn = picture(banner, CFG.get("logo_print_path") or CFG.get("logo_path"), height=LOGO_H)
    if not drawn:
        run_text(banner, CFG["company_name"], 13, COLORS["dark"], bold=True)
        drawn = pdfmetrics.stringWidth(CFG["company_name"], FONT_BB, 13)
    size, shown = fit(title, FONT_B, 15, content_w - drawn - 12, minimum=9)   # keep clear of the wordmark
    run_text(banner, "\t", size, COLORS["dark"])
    run_text(banner, shown, size, COLORS["dark"])
    if draft:
        head._element.append(parse_xml(WATERMARK.format(ns=nsdecls("w"), accent=COLORS["accent"],
                                                        family=escape(FAMILY))))

    foot = section.footer
    foot.is_linked_to_previous = False
    line = shape(plain(foot.paragraphs[0]), size=7.5, leading=11)
    line.paragraph_format.tab_stops.add_tab_stop(Pt(content_w), WD_TAB_ALIGNMENT.RIGHT)
    rule(line, "top", 0.5, COLORS["hairline"])
    if picture(line, CFG.get("logo_mark_path"), width=MARK_W):
        run_text(line, "   ", 7.5, COLORS["muted"])
    run_text(line, f"{clean(CFG.get('contact_line', ''))}\tPage ", 7.5, COLORS["muted"])
    line._p.append(parse_xml(PAGE_FIELD.format(ns=nsdecls("w"), family=escape(FAMILY, {'"': "&quot;"}),
                                               half=15, color=COLORS["muted"])))


def new_document(title, draft=False, landscape=False):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = FAMILY, Pt(DESC[0])
    normal.font.color.rgb = RGBColor.from_string(COLORS["body"])
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(DESC[1])
    section = document.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = (Pt(PAGE_H), Pt(PAGE_W)) if landscape else (Pt(PAGE_W), Pt(PAGE_H))
    section.left_margin = section.right_margin = Pt(M)
    section.top_margin, section.bottom_margin = Pt(PAGE_H - BODY_TOP), Pt(BODY_BOTTOM)
    section.header_distance, section.footer_distance = Pt(HEADER_TOP), Pt(FOOTER_BOTTOM)
    furniture(section, title, draft, content_width(document))
    for paragraph in list(document.paragraphs):       # the template's empty first line
        paragraph._p.getparent().remove(paragraph._p)
    return document


def content_width(document):
    section = document.sections[0]
    return (section.page_width - section.left_margin - section.right_margin) / 12700


def finish(document, pid, project, title, output_dir, rev=None, subject=""):
    """Document metadata, a closing paragraph Word requires after a table, and the file."""
    shape(document.add_paragraph(), size=1, leading=1)
    document.core_properties.title = f"{clean(project.get('Project Name')) or pid} - {title}"
    document.core_properties.author = CFG["company_name"]
    if rev:
        subject = f"Revision {rev} - exported {dt.date.today().isoformat()}"
    document.core_properties.subject = subject
    destination = Path(output_dir or OUT)
    destination.mkdir(exist_ok=True)
    return destination


def filename(pid, project, title, rev=None, parts=(), suffix=""):
    pieces = [re.sub(r"[^A-Za-z0-9_-]", "_", pid)[:50],
              re.sub(r"[^A-Za-z0-9]+", "_", clean(project.get("Project Name"))).strip("_")[:100],
              re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")[:60],
              f"R{rev}" if rev else "", *parts]
    return "_".join(piece for piece in pieces if piece) + suffix + ".docx"


# --------------------------------------------------------------------------
# Schedule blocks
# --------------------------------------------------------------------------
def band(document, values, content_w, controls=False):
    """Client / project / address / date fields under the header."""
    widths = [content_w / 4] * 4
    table = grid(document, widths, pad_x=0, pad_y=0)
    cells = add_row(table, widths, align=WD_ALIGN_VERTICAL.TOP)
    row_height(table.rows[0], BAND_H)
    keep_whole(table.rows[0])
    field_w = widths[0] - 2 * BAND_INSET
    for cell, (field, value) in zip(cells, values):
        shade(cell, COLORS["band"])
        margins(cell._tc.get_or_add_tcPr(), "w:tcMar", TCPR_ORDER,
                top=17, left=BAND_INSET, bottom=0, right=BAND_INSET)
        write(cell.paragraphs[0], field, 6.5, COLORS["muted"], bold=True, leading=8)
        line = shape(cell.add_paragraph(), size=9, leading=13, before=7)
        rule(line, "bottom", 0.5, COLORS["field_rule"])
        tag = field.lower().replace(" / ", "_").replace(" ", "_")
        # project_lot prints two workbook columns at once, so it is shown but never editable
        if controls and tag != "project_lot":
            control(line, tag, value, alias=field.title(), size=9, color=COLORS["text"])
        else:
            run_text(line, fit(value, FONT_B, 9, field_w)[1], 9, COLORS["text"])
    return table


def heading(document, text, content_w, first=False):
    """Letter-spaced section heading whose rule runs from the text past the right margin."""
    paragraph = shape(document.add_paragraph(), size=ITEM[0], leading=ITEM[1],
                      before=0 if first else SECTION_GAP, indent=PAD_X, keep=True)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Pt(content_w + HEADING_OVERHANG), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
    run_text(paragraph, f"{spaced_caps(text)} \t", ITEM[0], COLORS["dark"], bold=True)
    return paragraph


def item_cell(cell, row, hidden):
    """The item column: the item name, with its room beneath, minus whatever is hidden."""
    first = True
    for column, (size, leading), color in [("Item", ITEM, COLORS["dark"]), ("Room / Area", ROOM, COLORS["muted"])]:
        if column in hidden or not clean(row.get(column)):
            continue
        write(cell.paragraphs[0] if first else cell.add_paragraph(), row[column], size, color, leading=leading)
        first = False


def rows_table(document, rows, content_w, draft=False, controls=False, hidden=()):
    reserved = (STATUS_W if draft else 0) + (REMOVE_W if controls else 0)
    code_w, item_w, desc_w = row_widths(hidden, reserved)
    desc_w += content_w - CONTENT_W              # row_widths measures the portrait page
    widths = [width for width in (code_w, item_w) if width] + [desc_w]
    widths += [STATUS_W] if draft else []
    widths += [REMOVE_W] if controls else []
    table = grid(document, widths)
    for position, row in enumerate(rows):
        cells = iter(add_row(table, widths))
        keep_whole(table.rows[position])
        if code_w:
            write(next(cells).paragraphs[0], row.get("_code", ""), CODE[0], COLORS["muted"], leading=CODE[1])
        if item_w:
            item_cell(next(cells), row, hidden)
        body = shape(next(cells).paragraphs[0], size=DESC[0], leading=DESC[1])
        if controls:
            control(body, f"row{int(row['_row'])}_desc", description(row, hidden=hidden), alias="Description")
            checkbox(shape(next(cells).paragraphs[0], size=DESC[0], leading=DESC[1],
                           align=WD_ALIGN_PARAGRAPH.CENTER), f"row{int(row['_row'])}_remove")
        else:
            run_text(body, description(row, hidden=hidden), DESC[0], COLORS["body"])
            url = product_link(row, hidden)
            if url:
                run_text(body, " \u2013 ", DESC[0], COLORS["body"])
                hyperlink(body, url, PRODUCT_LINK, DESC[0], COLORS["accent"])
        if draft:
            write(next(cells).paragraphs[0], clean(row.get("Lookup Status")) or "Not run", STATUS[0],
                  COLORS["muted"], leading=STATUS[1], align=WD_ALIGN_PARAGRAPH.RIGHT)
        if position < len(rows) - 1:
            for cell in table.rows[position].cells:
                cell_rule(cell, "bottom", 0.4, COLORS["hairline"])
    return table


def detail_tables(document, project, rows, content_w, hidden=()):
    """Project and selection custom fields, in the same two-column rhythm."""
    groups = [] if "details" in hidden else presentation_details(project, rows)
    if not groups:
        return
    heading(document, "Additional details", content_w)
    widths = [CODE_W + ITEM_W, content_w - CODE_W - ITEM_W]
    for title, fields in groups:
        write(document.add_paragraph(), title, ITEM[0], COLORS["dark"], bold=True, leading=ITEM[1],
              before=8, after=3, indent=PAD_X, keep=True)
        table = grid(document, widths)
        for position, field in enumerate(fields):
            cells = add_row(table, widths, align=WD_ALIGN_VERTICAL.TOP)
            keep_whole(table.rows[position])
            write(cells[0].paragraphs[0], field["name"], ITEM[0], COLORS["dark"], leading=ITEM[1])
            write(cells[1].paragraphs[0], field["value"] or "Not provided", DESC[0], COLORS["body"],
                  leading=DESC[1])
            if position < len(fields) - 1:
                for cell in cells:
                    cell_rule(cell, "bottom", 0.4, COLORS["hairline"])


def sign_off(document):
    write(document.add_paragraph(), SCHEDULE.get("note"), NOTE[0], COLORS["muted"], leading=NOTE[1],
          before=SECTION_GAP, after=16, indent=PAD_X, keep=True)
    labels = list(SCHEDULE.get("signatures") or [])
    table = grid(document, SIGN_W, pad_x=0, pad_y=0, indent=SIGN_X)
    cells = add_row(table, SIGN_W, align=WD_ALIGN_VERTICAL.BOTTOM)
    row_height(table.rows[0], SIGN_H)
    keep_whole(table.rows[0])
    for cell, label in zip(cells, labels):
        margins(cell._tc.get_or_add_tcPr(), "w:tcMar", TCPR_ORDER, top=4, left=0, bottom=0, right=0)
        rule(shape(cell.paragraphs[0], size=7.5, leading=20), "bottom", 0.5, COLORS["dark"])
        write(cell.add_paragraph(), label, 7.5, COLORS["muted"], leading=11)


def schedule_body(document, project, rows, content_w, draft=False, controls=False, hidden=()):
    band(document, form_band_values(project) if controls else band_values(project), content_w, controls)
    shape(document.add_paragraph(), size=BAND_GAP, leading=BAND_GAP)
    if controls:
        write(document.add_paragraph(), "Tick REMOVE to delete a line when this document is imported.",
              6.5, COLORS["muted"], leading=9, align=WD_ALIGN_PARAGRAPH.RIGHT, after=4)
    for index, (section, items) in enumerate(group_sections(rows, hidden).items()):
        if section:
            heading(document, section, content_w, first=index == 0 and not controls)
        rows_table(document, items, content_w, draft, controls, hidden)


# --------------------------------------------------------------------------
# Documents
# --------------------------------------------------------------------------
def build(pid, project, rows, draft=False, output_dir=None, title=None, prefix=None, rev=None, hidden=()):
    """The client finish schedule, editable in Word."""
    title, rows = schedule_title(title), numbered(rows, code_prefix(prefix))
    document = new_document(title, draft)
    width = content_width(document)
    schedule_body(document, project, rows, width, draft=draft, hidden=hidden)
    detail_tables(document, project, rows, width, hidden)
    sign_off(document)
    destination = finish(document, pid, project, title, output_dir, rev)
    out = destination / filename(pid, project, title, rev, suffix="_DRAFT" if draft else "")
    document.save(out)
    return out


def build_fillable(pid, project, rows, output_dir=None, title=None, prefix=None, hidden=()):
    """The working copy: client details, descriptions and remove ticks as content controls."""
    title, rows = schedule_title(title), numbered(rows, code_prefix(prefix))
    document = new_document(title)
    width = content_width(document)
    schedule_body(document, project, rows, width, controls=True, hidden=hidden)
    sign_off(document)
    # Which columns the copy was built without, so re-importing it cannot mistake a shortened
    # description for one the client reworded by hand. Carried invisibly, like the PDF's field.
    control(shape(document.add_paragraph(), size=1, leading=1), HIDDEN_FIELD,
            ",".join(sorted(hidden)), alias="Hidden columns", size=1, vanish=True)
    destination = finish(document, pid, project, title, output_dir,
                         subject=f"Fillable working copy for {pid}")
    out = destination / filename(pid, project, title, parts=("FILLABLE",))
    document.save(out)
    return out


def build_spec(pid, project, rows, output_dir=None, hidden=()):
    """The internal spec sheet: every priced line item, landscape, grouped by section."""
    document = new_document(SPEC_TITLE, landscape=True)
    width = content_width(document)
    band(document, band_values(project), width)
    shape(document.add_paragraph(), size=BAND_GAP, leading=BAND_GAP)
    columns = [(column, share) for column, share, key in SPEC_COLUMNS if key not in hidden]
    if not columns:
        columns = [("Item", 1)]
    widths = [width * share / sum(share for _, share in columns) for _, share in columns]
    for index, (section, items) in enumerate(group_sections(rows, hidden).items()):
        if section:
            heading(document, section, width, first=index == 0)
        table = grid(document, widths, pad_y=4)
        titles = add_row(table, widths, align=WD_ALIGN_VERTICAL.BOTTOM)
        keep_whole(table.rows[0], header=True)
        for cell, (column, _) in zip(titles, columns):
            write(cell.paragraphs[0], column.upper(), 6.5, COLORS["muted"], bold=True, leading=9)
            cell_rule(cell, "bottom", 0.5, COLORS["heading_rule"])
        for position, row in enumerate(items):
            cells = add_row(table, widths, align=WD_ALIGN_VERTICAL.TOP)
            keep_whole(table.rows[position + 1])
            for cell, (column, _) in zip(cells, columns):
                url = product_link(row, hidden) if column == PRODUCT_LINK else ""
                if url:
                    hyperlink(shape(cell.paragraphs[0], size=8, leading=10.5), url, PRODUCT_LINK,
                              8, COLORS["accent"])
                elif column != PRODUCT_LINK:
                    write(cell.paragraphs[0], row.get(column, ""), 8, COLORS["body"], leading=10.5)
                if position < len(items) - 1:
                    cell_rule(cell, "bottom", 0.4, COLORS["hairline"])
            if "Include in Lookbook" not in hidden and clean(row.get("Include in Lookbook")).lower() == "no":
                write(cells[0].add_paragraph(), "Not in lookbook", 6.5, COLORS["muted"], leading=9)
    destination = finish(document, pid, project, SPEC_TITLE, output_dir,
                         subject=f"Internal spec sheet for {pid}")
    out = destination / filename(pid, project, SPEC_TITLE)
    document.save(out)
    return out


# --------------------------------------------------------------------------
# Reading a filled copy back
# --------------------------------------------------------------------------
def form_text(value):
    """Word stores what was typed, so nothing is transliterated the way the PDF form is."""
    return clean(value)


def control_text(sdt):
    content = sdt.find(qn("w:sdtContent"))
    return "".join(node.text or "" for node in content.iter(qn("w:t"))) if content is not None else ""


def ticked(properties, sdt):
    """A Word checkbox, or the glyph left behind by an editor that dropped the control."""
    checked = properties.find(f'{qn("w14:checkbox")}/{qn("w14:checked")}')
    if checked is not None:
        return clean(checked.get(qn("w14:val"))).lower() in {"1", "true", "on"}
    return clean(control_text(sdt)).lower() in {"\u2612", "x", "yes"}


def read(path_or_stream):
    """Field values from a filled copy: client details, per-row text, remove ticks, hidden columns."""
    document = Document(path_or_stream)
    details, lines, hidden = {}, {}, set()
    for sdt in document.element.body.iter(qn("w:sdt")):
        properties = sdt.find(qn("w:sdtPr"))
        tag = properties.find(qn("w:tag")) if properties is not None else None
        name = clean(tag.get(qn("w:val"))) if tag is not None else ""
        match = re.fullmatch(r"row(\d+)_(desc|remove)", name)
        if match:
            entry = lines.setdefault(int(match[1]), {})
            if match[2] == "desc":
                entry["description"] = clean(control_text(sdt))
            else:
                entry["remove"] = ticked(properties, sdt)
        elif name in {"client", "address", "date"}:
            details[name] = clean(control_text(sdt))
        elif name == HIDDEN_FIELD:
            # a copy from an older build carries no marker, which reads as "nothing was hidden"
            hidden = {column for column in control_text(sdt).split(",") if column in HIDEABLE}
    return details, lines, hidden


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--project", required=True, help="Project ID, e.g. UH-101")
    parser.add_argument("--draft", action="store_true", help="Internal review copy with status labels")
    parser.add_argument("--fillable", action="store_true", help="Working copy with content controls")
    parser.add_argument("--spec", action="store_true", help="Internal priced spec sheet")
    parser.add_argument("--verified-only", action="store_true", help="Only include rows marked Verified")
    parser.add_argument("--title", help=f"Schedule title (default: {SCHEDULE.get('title', 'Selections')})")
    parser.add_argument("--prefix", help=f"Item code prefix (default: {SCHEDULE.get('code_prefix', 'EX')})")
    parser.add_argument("--hide", default="", help=f"Columns to leave out, comma separated: {', '.join(HIDEABLE)}")
    args = parser.parse_args()
    try:
        hidden = hidden_columns(args.hide)
    except ValueError as error:
        parser.error(str(error))
    projects, items = load_data(args.verified_only)
    if args.project not in projects or not items.get(args.project):
        parser.error("Project not found or no selections match the export options.")
    project, rows = projects[args.project], items[args.project]
    if args.spec:
        print(build_spec(args.project, project, rows, hidden=hidden))
    elif args.fillable:
        print(build_fillable(args.project, project, rows, title=args.title, prefix=args.prefix, hidden=hidden))
    else:
        if not args.draft and any(row["Lookup Status"] != "Verified" for row in rows):
            parser.error("Review all selections first, or use --draft or --verified-only.")
        print(build(args.project, project, rows, args.draft, title=args.title, prefix=args.prefix,
                    hidden=hidden))


if __name__ == "__main__":
    main()
