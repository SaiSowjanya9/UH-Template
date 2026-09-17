import argparse
import io
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from build_lookbook import CFG, OUT, fmt_date, get_image, group_sections, load_data
from branding import load_logo


COLORS = {key: value.lstrip("#") for key, value in CFG["colors"].items()}


def rectangle(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def text(slide, x, y, w, h, value, size=14, color=None, serif=False, bold=False, link=None, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, line in enumerate(str(value or "").split("\n")):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if align is not None:
            p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = "Georgia" if serif else "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color or COLORS["text"])
        if link:
            run.hyperlink.address = link
    return box


def picture(slide, src, x, y, w, h, brand=False):
    image = load_logo(src) if brand else get_image(src)
    if image is None:
        return False
    if not brand:
        image = Image.frombytes("RGB", image.getSize(), image.getRGBData())
    iw, ih = image.size
    stream = io.BytesIO()
    image.save(stream, "PNG")
    stream.seek(0)
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(stream, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), width=Inches(pw), height=Inches(ph))
    return True


def new_slide(prs, project, draft, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(COLORS["dark"] if dark else COLORS["paper"])
    color = COLORS["paper"] if dark else COLORS["muted"]
    has_mark = picture(slide, CFG.get("logo_mark_path"), .55, 7.02, .24, .24, brand=True)
    text(slide, .88 if has_mark else .55, 7.1, 2.3, .2, CFG["company_name"].upper(), 8, color, bold=True)
    text(slide, 3.2, 7.1, 6.8, .2, project.get("Project Name", ""), 8, color, align=PP_ALIGN.CENTER)
    text(slide, 11.95, 7.1, .8, .2, f"{len(prs.slides):02}", 8, color, align=PP_ALIGN.RIGHT)
    if draft:
        text(slide, 10.45, .24, 2.3, .25, "DRAFT · FOR REVIEW", 9, COLORS["accent"], bold=True, align=PP_ALIGN.RIGHT)
    return slide


def title(slide, eyebrow, heading):
    text(slide, .55, .42, 10, .25, eyebrow.upper(), 10, COLORS["accent"], bold=True)
    text(slide, .55, .87, 12.1, .6, heading, 30, serif=True)


def product_card(slide, row, x, draft):
    rectangle(slide, x, 1.7, 5.95, 5.05, "FFFFFF")
    if not picture(slide, row.get("Image URL"), x + .2, 1.85, 5.55, 1.85):
        rectangle(slide, x + .18, 1.85, 5.59, 1.85, COLORS["paper"])
        text(slide, x + .4, 2.6, 5.1, .3, "PRODUCT IMAGE PENDING", 10, COLORS["muted"], align=PP_ALIGN.CENTER)
    text(slide, x + .22, 3.88, 5.5, .3, " / ".join(v for v in [row.get("Room / Area"), row.get("Item")] if v), 10, COLORS["accent"], bold=True)
    text(slide, x + .22, 4.25, 5.5, .62, row.get("Product Name") or row.get("Item"), 20, serif=True)
    details = [" · ".join(v for v in [row.get("Manufacturer"), f"Model {row['Model #']}" if row.get("Model #") else ""] if v),
               " · ".join(v for v in [f"Finish: {row['Finish / Color']}" if row.get("Finish / Color") else "", f"Qty: {row['Qty']}" if row.get("Qty") else ""] if v)]
    if row.get("Client Notes"):
        details.append(row["Client Notes"])
    text(slide, x + .22, 4.99, 5.5, 1.15, "\n".join(filter(None, details)), 12, COLORS["muted"])
    url = row.get("Product URL", "")
    if url.startswith(("https://", "http://")):
        text(slide, x + .22, 6.34, 2.35, .22, "VIEW PRODUCT →", 10, COLORS["accent"], bold=True, link=url)
    if draft:
        text(slide, x + 2.6, 6.34, 3.1, .22, row.get("Lookup Status", "Not run"), 9, COLORS["muted"], align=PP_ALIGN.RIGHT)


def build(pid, project, rows, draft=False, output_dir=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = f"{project.get('Project Name', pid)} — Selections Lookbook"
    prs.core_properties.author = CFG["company_name"]
    sections = group_sections(rows)
    slide = new_slide(prs, project, draft, dark=True)
    rectangle(slide, 7.25, 0, 6.083, 6.95, COLORS["paper"])
    if not picture(slide, CFG.get("logo_path"), .6, .6, 4.2, .65, brand=True):
        text(slide, .6, .65, 5.9, .5, CFG["company_name"].upper(), 22, "FFFFFF", bold=True)
    text(slide, .6, 2, 5.9, .3, CFG.get("tagline", "Selections Lookbook").upper(), 11, COLORS["accent"], bold=True)
    text(slide, .6, 2.55, 5.9, 1.8, project.get("Project Name", ""), 40, "FFFFFF", serif=True)
    details = [f"Prepared for {project['Client Name']}" if project.get("Client Name") else "", project.get("Address", ""),
               project.get("Plan / Elevation", ""), fmt_date(project.get("Presentation Date", ""))]
    text(slide, .6, 4.75, 5.9, 1.5, "\n".join(filter(None, details)), 14, "FFFFFF")
    if not picture(slide, project.get("Cover Image"), 7.55, .5, 5.5, 5.7):
        if not picture(slide, CFG.get("logo_mark_path"), 8.75, 1.75, 3, 2.1, brand=True):
            text(slide, 7.75, 2.35, 5, 1.5, "UH", 100, COLORS["accent"], serif=True, align=PP_ALIGN.CENTER)
        text(slide, 7.75, 4.1, 5, .45, f"{len(sections)} CATEGORIES  /  {len(rows)} SELECTIONS", 12, COLORS["muted"], align=PP_ALIGN.CENTER)
    section_list = list(sections.items())
    for start in range(0, len(section_list), 8):
        slide = new_slide(prs, project, draft)
        title(slide, "The details make the home", "Your selections, considered.")
        for i, (section, items) in enumerate(section_list[start:start + 8], start=start + 1):
            y = 1.9 + (i - start - 1) * .55
            text(slide, .65, y, .6, .35, f"{i:02}", 18, COLORS["accent"], serif=True)
            text(slide, 1.4, y, 8.3, .4, section, 20, serif=True)
            text(slide, 10.1, y + .05, 2.1, .3, f"{len(items)} selections", 12, COLORS["muted"], align=PP_ALIGN.RIGHT)
    for index, (section, items) in enumerate(sections.items(), start=1):
        for offset in range(0, len(items), 2):
            slide = new_slide(prs, project, draft)
            title(slide, f"{index:02} / Material & finish selections", section + (" — continued" if offset else ""))
            for position, row in enumerate(items[offset:offset + 2]):
                product_card(slide, row, .55 + position * 6.28, draft)
    for offset in range(0, len(rows), 8):
        slide = new_slide(prs, project, draft)
        title(slide, "Reference", "Selections schedule" + (" — continued" if offset else ""))
        chunk = rows[offset:offset + 8]
        headings = ["Section / Room", "Item", "Manufacturer / Model", "Finish / Qty", "Product"]
        table = slide.shapes.add_table(len(chunk) + 1, 5, Inches(.55), Inches(1.8), Inches(12.2), Inches(.5 * (len(chunk) + 1))).table
        for j, width in enumerate([2.5, 2.3, 3.15, 2.8, 1.45]):
            table.columns[j].width = Inches(width)
        for i in range(len(chunk) + 1):
            row = chunk[i - 1] if i else None
            values = headings if i == 0 else ["\n".join(filter(None, [row["Section"], row["Room / Area"]])), row["Item"],
                                                   "\n".join(filter(None, [row["Manufacturer"], row["Model #"]])),
                                                   "\n".join(filter(None, [row["Finish / Color"], f"Qty: {row['Qty']}" if row["Qty"] else ""])),
                                                   "View product" if row["Product URL"].startswith(("http://", "https://")) else "Pending"]
            for j, value in enumerate(values):
                cell = table.cell(i, j)
                cell.text = value
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(COLORS["dark"] if i == 0 else ("FFFFFF" if i % 2 else COLORS["paper"]))
                cell.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.name, run.font.size = "Aptos", Pt(10)
                        run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else COLORS["text"])
                        if i and j == 4 and value == "View product":
                            run.hyperlink.address = row["Product URL"]
    slide = new_slide(prs, project, draft, dark=True)
    if not picture(slide, CFG.get("logo_path"), 4.15, 2.1, 5.03, .75, brand=True):
        text(slide, 1, 2.2, 11.3, .6, CFG["company_name"].upper(), 24, "FFFFFF", bold=True, align=PP_ALIGN.CENTER)
    text(slide, 1, 3.15, 11.3, .8, "A home, thoughtfully yours.", 36, "FFFFFF", serif=True, align=PP_ALIGN.CENTER)
    text(slide, 1, 4.25, 11.3, .6, CFG.get("contact_line", ""), 12, COLORS["paper"], align=PP_ALIGN.CENTER)
    text(slide, 2.1, 5.9, 9.1, .7, CFG.get("disclaimer", ""), 9, COLORS["muted"], align=PP_ALIGN.CENTER)
    destination = Path(output_dir or OUT)
    destination.mkdir(exist_ok=True)
    safe_pid = re.sub(r"[^A-Za-z0-9_-]", "_", pid)[:50]
    name = re.sub(r"[^A-Za-z0-9]+", "_", project.get("Project Name", "")).strip("_")[:100]
    output = destination / f"{safe_pid}_{name}_Lookbook{'_DRAFT' if draft else ''}.pptx"
    prs.save(output)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate an editable UH Homes PowerPoint from the master workbook.")
    parser.add_argument("--project", required=True)
    parser.add_argument("--draft", action="store_true")
    parser.add_argument("--verified-only", action="store_true")
    args = parser.parse_args()
    projects, items = load_data(args.verified_only)
    if args.project not in projects or not items.get(args.project):
        parser.error("Project not found or no selections match the export options.")
    rows = items[args.project]
    if not args.draft and any(row["Lookup Status"] != "Verified" for row in rows):
        parser.error("Review all selections first, or use --draft or --verified-only.")
    print(build(args.project, projects[args.project], rows, args.draft))
