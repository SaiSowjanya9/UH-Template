"""
Bring the master workbook up to date for direct data entry in Excel.

Usage:
    python migrate_workbook.py            # update the workbook named in .env / the default
    python migrate_workbook.py --dry-run  # report what would change

Safe to re-run: it only fixes presentation and guidance. Row data is never rewritten.
Saving goes through WorkbookStore, so the workbook is validated and backed up first, and
the save is refused while Excel holds the file open.
"""
import argparse

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from common import WORKBOOK, Sheet, clean
from workbook_store import (CLIENT_STATUSES, SELECTION_FIELDS, STATUSES, WorkbookStore,
                            ensure_list_values, ensure_selection_columns, ensure_validations)

# Columns the user fills in by hand versus those the lookup maintains.
ENTRY_COLUMNS = ["Project ID", "Section", "Room / Area", "Item", "Manufacturer", "Model #",
                 "Finish / Color", "Qty", "Unit Price", "Client Notes", "Include in Lookbook", "Client Status"]
WIDTHS = {"Selections": {"Image URL": 32, "Product Name": 30, "Lookup Notes": 34, "Client Notes": 28,
                         "Include in Lookbook": 17, "Client Status": 15, "Checked On": 12, "Unit Price": 12},
          "Projects": {"Cover Image": 30, "Presentation Date": 18}}

INSTRUCTIONS = [
    ("UH Homes - Selections Tracker", "title"),
    ("", ""),
    ("This workbook is the source of truth. You can type directly into it, or use the web app "
     "(python app.py, then open the address it prints). Both read and write these same sheets.", ""),
    ("", ""),
    ("How the workflow runs", "heading"),
    ("1. Projects sheet - one row per project, with a unique Project ID such as UH-101.", ""),
    ("2. Selections sheet - one row per product. Fill the entry columns listed below.", ""),
    ("3. Product links - the web app finds them automatically while it is running, or run "
     "python find_urls.py (add --project UH-101 for one project, --dry-run to preview).", ""),
    ("4. Review each link against the real product page, then set Lookup Status to Verified. "
     "Verified rows are never overwritten automatically.", ""),
    ("5. Client Status records the client's decision: Proposed, Presented, Approved, Rejected, Changed.", ""),
    ("6. Export the finish schedule from the web app, or run "
     "python build_lookbook.py --project UH-101 (PDF) or build_powerpoint.py (editable deck).", ""),
    ("", ""),
    ("Columns you fill in (Selections)", "heading"),
    ("Project ID       Must match a Project ID on the Projects sheet. Dropdown.", ""),
    ("Section          Category and print order; values come from the Lists sheet. Dropdown.", ""),
    ("Room / Area      Optional location, for example Primary Bath.", ""),
    ("Item             What the client sees, for example Kitchen Faucet. Required.", ""),
    ("Manufacturer     Brand name. Add its official domain on the Manufacturers sheet.", ""),
    ("Model #          Manufacturer model or SKU. Needed for automatic product lookup.", ""),
    ("Finish / Color   For example Matte Black or SW 7010.", ""),
    ("Qty              Leave blank or 1 to hide the quantity from the client schedule.", ""),
    ("Unit Price       Cost per unit. The Price Schedule sheet multiplies it by Qty and totals it.", ""),
    ("Client Notes     Short note printed with the item on the client schedule.", ""),
    ("Include in Lookbook   No hides the row from client exports without deleting it. Dropdown.", ""),
    ("Client Status    Proposed / Presented / Approved / Rejected / Changed. Dropdown.", ""),
    ("", ""),
    ("Price Schedule sheet", "heading"),
    ("Every selection grouped by project and category with S.No, Code, Description, Qty, Unit Price "
     "and Price, plus a subtotal per category, a total per project and a grand total.", ""),
    ("It is rebuilt automatically each time the workbook is saved, so do not type into it. "
     "Enter costs in the Unit Price column on Selections and they flow through.", ""),
    ("Prices stay internal: they never appear on the client PDF or PowerPoint.", ""),
    ("", ""),
    ("Columns filled in for you", "heading"),
    ("Product URL, Product Name, Image URL, Lookup Status, Checked On, Lookup Notes are written by the "
     "product lookup. You may correct any of them by hand.", ""),
    ("On the Projects sheet, Total Items, Links Found, Verified and Needs Review are formulas.", ""),
    ("", ""),
    ("Rules worth knowing", "heading"),
    ("- Changing Item, Product Name, Manufacturer, Model # or Finish clears the old link and "
     "verification, because they no longer describe the same product.", ""),
    ("- A Verified link older than six months is flagged for re-checking in the web app.", ""),
    ("- Close this workbook in Excel before saving from the web app; the app will not write "
     "to a file Excel has open.", ""),
    ("- Every save keeps a timestamped copy in the backups folder.", ""),
    ("- Product URLs and images are internal review aids and never appear on client documents.", ""),
    ("- Section order on the Lists sheet controls the order of categories in the client schedule.", ""),
]


def write_instructions(ws):
    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, len(INSTRUCTIONS))):
        for cell in row:
            cell.value = None
    for index, (text, kind) in enumerate(INSTRUCTIONS, start=1):
        cell = ws.cell(index, 1, text or None)
        cell.font = Font(name="Calibri", size=16 if kind == "title" else 11, bold=kind in {"title", "heading"})
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 118
    ws.sheet_view.showGridLines = False


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--dry-run", action="store_true", help="Report the changes without saving")
    args = parser.parse_args()

    store = WorkbookStore(WORKBOOK)
    wb, revision = store.snapshot()
    changes = []

    sheet = ensure_selection_columns(wb)
    missing = [field for field in SELECTION_FIELDS if field not in sheet.cols]
    if missing:
        raise SystemExit(f"Selections is missing required columns: {', '.join(missing)}")
    ensure_list_values(wb)
    ensure_validations(wb)
    changes.append(f"Lists: Lookup Status offers all {len(STATUSES)} values, Client Status offers {len(CLIENT_STATUSES)}")
    changes.append("Selections: dropdowns present for Project ID, Section, Lookup Status, Include in Lookbook, Client Status")

    write_instructions(wb["Instructions"])
    changes.append("Instructions: rewritten to cover the web app, every column, and the current rules")

    selections = wb["Selections"]
    if selections.freeze_panes != "A2":
        selections.freeze_panes = "A2"   # keep only the header row in view
        changes.append("Selections: froze the header row only")
    for name in ["Projects", "Manufacturers", "Selections"]:
        ws = wb[name]
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        if ws.auto_filter.ref != ref:
            ws.auto_filter.ref = ref
            changes.append(f"{name}: filter covers every column")
    for name, widths in WIDTHS.items():
        columns = Sheet(wb[name]).cols
        for header, width in widths.items():
            letter = get_column_letter(columns[header])
            if (wb[name].column_dimensions[letter].width or 0) < width:
                wb[name].column_dimensions[letter].width = width
                changes.append(f"{name}: widened {header}")

    entry = [header for header in ENTRY_COLUMNS if header in sheet.cols]
    print("Entry columns on Selections:", ", ".join(entry))
    for line in changes:
        print(" -", line)
    if args.dry_run:
        print(f"Dry run - {len(changes)} change(s), nothing saved.")
        return
    store.save(wb, revision, "workbook entry migration")
    print(f"Saved. {len(changes)} change(s) applied; a backup was kept.")


if __name__ == "__main__":
    main()
