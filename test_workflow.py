import io
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import common
import find_urls
import build_powerpoint


class MatchingTests(unittest.TestCase):
    def test_model_boundaries(self):
        self.assertTrue(find_urls.exact_model("Model K-596-VS", "k596vs"))
        self.assertTrue(find_urls.exact_model("Model K596VS", "k596vs"))
        self.assertFalse(find_urls.exact_model("Model K-596-VS2", "k596vs"))
        self.assertFalse(find_urls.exact_model("Model 17594ESRS", "7594esrs"))

    def test_category_page_is_not_product_match(self):
        html = '<html><title>All faucets</title><body>K-596-VS and many more</body></html>'
        with patch("find_urls.fetch_html", return_value=(html, "https://example.com/category")):
            self.assertEqual(find_urls.inspect_page("https://example.com/category", "k596vs"), (False, "", ""))

    def test_structured_category_is_not_a_product_page(self):
        html = '<script type="application/ld+json">[{"@type":"Product","sku":"K-596-VS"},{"@type":"Product","sku":"K-597-VS"}]</script>'
        with patch("find_urls.fetch_html", return_value=(html, "https://example.com/category")):
            self.assertEqual(find_urls.inspect_page("https://example.com/category", "k596vs"), (False, "", ""))

    def test_model_does_not_match_hyphenated_variant_or_prefix(self):
        self.assertFalse(find_urls.exact_model("K-596-VS-2", "k596vs"))
        self.assertFalse(find_urls.exact_model("PREFIX-K-596-VS", "k596vs"))

    def test_structured_product_match(self):
        html = '<script type="application/ld+json">{"@type":"Product","sku":"K-596-VS","name":"Kitchen faucet","image":"/faucet.jpg"}</script>'
        with patch("find_urls.fetch_html", return_value=(html, "https://example.com/faucet")):
            self.assertEqual(find_urls.inspect_page("https://example.com/faucet", "k596vs"),
                             (True, "Kitchen faucet", "https://example.com/faucet.jpg"))

    def test_fallback_after_unconfirmed_official_hit(self):
        queries = []
        def search(q):
            queries.append(q)
            return [{"url": "https://brand.example/category" if len(queries) == 1 else "https://shop.example/product", "title": "", "snippet": ""}]
        def inspect(url, model):
            return ("shop.example" in url, "Product", "")
        with patch("find_urls.inspect_page", side_effect=inspect):
            result = find_urls.lookup(search, "Brand", "ABC123", "brand.example")
        self.assertEqual(len(queries), 2)
        self.assertEqual(result["status"], "Found - retailer")

    def test_provider_errors_do_not_expose_keys(self):
        def search(q):
            raise RuntimeError("https://provider.example?api_key=secret-value")
        result = find_urls.lookup(search, "Brand", "ABC123", "")
        self.assertNotIn("secret-value", result["notes"])
        self.assertEqual(result["status"], "Search error")


class AppTests(unittest.TestCase):
    def setUp(self):
        import app
        from workbook_store import WorkbookStore
        self.module = app
        self.folder = tempfile.TemporaryDirectory()
        self.path = Path(self.folder.name) / "master.xlsx"
        shutil.copy2(common.WORKBOOK, self.path)
        self.store = WorkbookStore(self.path)
        self.patcher = patch.object(app, "store", self.store)
        self.patcher.start()
        self.client = app.app.test_client()
        self.headers = {"X-UH-Token": app.TOKEN}

    def tearDown(self):
        self.patcher.stop()
        self.folder.cleanup()

    def state(self):
        response = self.client.get("/api/state")
        self.assertEqual(response.status_code, 200, response.json)
        return response.json

    def post(self, path, data):
        return self.client.post(path, json={"revision": self.state()["revision"], **data}, headers=self.headers)

    def test_state_is_project_scoped_by_ids_and_contains_no_keys(self):
        state = self.state()
        ids = {p["Project ID"] for p in state["projects"]}
        self.assertTrue(ids)
        self.assertTrue(all(r["Project ID"] in ids for r in state["selections"]))
        self.assertNotIn("token", state)
        self.assertEqual(set(state["provider"]), {"name", "ready", "key_name"})

    def test_create_project_and_selection_preserve_text_models(self):
        response = self.post("/api/projects", {"create": True, "values": {"Project ID": "TEST-1", "Project Name": "Test House"}})
        self.assertEqual(response.status_code, 200, response.json)
        response = self.post("/api/selections", {"values": {"Project ID": "TEST-1", "Item": "Faucet", "Manufacturer": "Brand", "Model #": "00123"}})
        self.assertEqual(response.status_code, 200, response.json)
        record = next(r for r in self.state()["selections"] if r["Project ID"] == "TEST-1")
        self.assertEqual(record["Model #"], "00123")
        self.assertEqual(record["Lookup Status"], "Not run")
        self.assertEqual(len(list((self.path.parent / "backups").glob("*.xlsx"))), 2)

    def test_stale_revision_does_not_overwrite(self):
        revision = self.state()["revision"]
        self.post("/api/projects", {"create": True, "values": {"Project ID": "TEST-1", "Project Name": "Test House"}})
        before = self.path.read_bytes()
        response = self.client.post("/api/projects", json={"revision": revision, "values": {"Project ID": "TEST-2", "Project Name": "Stale House"}}, headers=self.headers)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.path.read_bytes(), before)

    def test_product_identity_change_clears_stale_metadata(self):
        row = self.state()["selections"][0]
        response = self.post("/api/selections", {"row": row["_row"], "values": {"Product URL": "https://example.com/product", "Product Name": "Previous product", "Image URL": "https://example.com/image.jpg"}})
        self.assertEqual(response.status_code, 200, response.json)
        response = self.post("/api/selections", {"row": row["_row"], "values": {"Lookup Status": "Verified"}, "confirm_verified": True})
        self.assertEqual(response.status_code, 200, response.json)
        response = self.post("/api/selections", {"row": row["_row"], "values": {"Model #": "NEW-MODEL"}})
        self.assertEqual(response.status_code, 200, response.json)
        current = next(r for r in self.state()["selections"] if r["_row"] == row["_row"])
        for key in ["Product URL", "Product Name", "Image URL", "Checked On"]:
            self.assertEqual(current[key], "")
        self.assertEqual(current["Lookup Status"], "Not run")

    def test_verification_requires_link_and_confirmation(self):
        row = self.state()["selections"][0]["_row"]
        self.assertEqual(self.post("/api/selections", {"row": row, "values": {"Lookup Status": "Verified"}}).status_code, 400)
        self.assertEqual(self.post("/api/selections", {"row": row, "values": {"Lookup Status": "Verified"}, "confirm_verified": True}).status_code, 400)

    def test_formula_like_input_is_written_as_literal_text(self):
        row = self.state()["selections"][0]["_row"]
        response = self.post("/api/selections", {"row": row, "values": {"Client Notes": "=1+1"}})
        self.assertEqual(response.status_code, 200, response.json)
        wb, _ = self.store.snapshot()
        sheet = common.Sheet(wb["Selections"])
        self.assertEqual(sheet.ws.cell(row, sheet.cols["Client Notes"]).data_type, "s")

    def test_missing_api_key_is_actionable(self):
        row = self.state()["selections"][0]["_row"]
        with patch.dict("os.environ", {"SEARCH_PROVIDER": "brave", "BRAVE_API_KEY": ""}):
            response = self.post(f"/api/selections/{row}/lookup", {})
        self.assertEqual(response.status_code, 400)
        self.assertIn("BRAVE_API_KEY", response.json["error"])

    def test_lookup_writes_candidate_but_not_verified(self):
        row = self.state()["selections"][0]["_row"]
        result = {"url": "https://example.com/product", "name": "Candidate", "image": "", "status": "Found - verify", "notes": ""}
        with patch.dict("os.environ", {"SEARCH_PROVIDER": "brave", "BRAVE_API_KEY": "test-only"}), patch("find_urls.lookup", return_value=result):
            response = self.post(f"/api/selections/{row}/lookup", {})
        self.assertEqual(response.status_code, 200, response.json)
        current = next(r for r in self.state()["selections"] if r["_row"] == row)
        self.assertEqual(current["Lookup Status"], "Found - verify")
        self.assertEqual(current["Product URL"], result["url"])

    def test_verified_rows_are_not_looked_up(self):
        row = self.state()["selections"][0]["_row"]
        self.post("/api/selections", {"row": row, "values": {"Product URL": "https://example.com/product"}})
        self.post("/api/selections", {"row": row, "values": {"Lookup Status": "Verified"}, "confirm_verified": True})
        with patch.dict("os.environ", {"SEARCH_PROVIDER": "brave", "BRAVE_API_KEY": "test-only"}), patch("find_urls.lookup") as lookup:
            response = self.post(f"/api/selections/{row}/lookup", {})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json["skipped"])
        lookup.assert_not_called()

    def test_final_export_rejects_unverified_selections(self):
        project = self.state()["projects"][0]["Project ID"]
        response = self.post("/api/presentation", {"project": project, "format": "pdf", "mode": "final"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("not verified", response.json["error"])

    def test_pdf_and_powerpoint_exports(self):
        from pptx import Presentation
        data = self.state()
        project = data["projects"][0]
        other = data["projects"][1]
        with patch("build_lookbook.get_image", return_value=None), patch("build_powerpoint.get_image", return_value=None):
            pdf = self.post("/api/presentation", {"project": project["Project ID"], "format": "pdf", "mode": "draft"})
            pptx = self.post("/api/presentation", {"project": project["Project ID"], "format": "pptx", "mode": "draft"})
        self.assertEqual(pdf.status_code, 200, pdf.json if pdf.is_json else "")
        self.assertTrue(pdf.data.startswith(b"%PDF-"))
        self.assertGreater(len(pdf.data), 5000)
        try:
            from pypdf import PdfReader
        except ImportError:
            PdfReader = None
        if PdfReader:
            reader = PdfReader(io.BytesIO(pdf.data))
            text = "\n".join(page.extract_text() for page in reader.pages)
            self.assertIn(project["Project Name"], text)
            self.assertNotIn(other["Project Name"], text)
            self.assertIn("DRAFT", text)
        self.assertEqual(pptx.status_code, 200, pptx.json if pptx.is_json else "")
        presentation = Presentation(io.BytesIO(pptx.data))
        text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn(project["Project Name"], text)
        self.assertNotIn(other["Project Name"], text)
        self.assertIn("DRAFT", text)
        self.assertGreater(len(presentation.slides), 4)

    def test_verified_only_export_omits_pending_and_hidden_rows(self):
        from pptx import Presentation
        data = self.state()
        row = data["selections"][0]
        self.post("/api/selections", {"row": row["_row"], "values": {"Product URL": "https://example.com/product"}})
        self.post("/api/selections", {"row": row["_row"], "values": {"Lookup Status": "Verified"}, "confirm_verified": True})
        with patch("build_powerpoint.get_image", return_value=None):
            response = self.post("/api/presentation", {"project": row["Project ID"], "format": "pptx", "mode": "verified"})
        self.assertEqual(response.status_code, 200, response.json if response.is_json else "")
        prs = Presentation(io.BytesIO(response.data))
        text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn("1 SELECTIONS", text)
        self.assertNotIn("DRAFT", text)
        self.post("/api/selections", {"row": row["_row"], "values": {"Include in Lookbook": "No"}})
        response = self.post("/api/presentation", {"project": row["Project ID"], "format": "pptx", "mode": "verified"})
        self.assertEqual(response.status_code, 400)

    def test_import_validates_before_replacement(self):
        before = self.path.read_bytes()
        response = self.client.post("/api/workbook/import", data={"file": (io.BytesIO(b"not an Excel file"), "broken.xlsx"), "revision": self.state()["revision"], "confirm": "replace-with-backup"}, headers=self.headers)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertFalse((self.path.parent / "backups").exists())

    def test_import_export_roundtrip(self):
        original = self.state()
        exported = self.client.get("/api/workbook")
        self.assertEqual(exported.status_code, 200)
        response = self.client.post("/api/workbook/import", data={"file": (io.BytesIO(exported.data), "tracker.xlsx"), "revision": original["revision"], "confirm": "replace-with-backup"}, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.state()["selections"], original["selections"])
        self.assertTrue(list((self.path.parent / "backups").glob("*.xlsx")))

    def test_invalid_project_and_missing_fields_are_rejected(self):
        self.assertEqual(self.post("/api/projects", {"values": {"Project ID": "../outside", "Project Name": "Invalid"}}).status_code, 400)
        self.assertEqual(self.post("/api/selections", {"values": {"Project ID": "UNKNOWN", "Item": "Item"}}).status_code, 400)
        self.assertEqual(self.post("/api/selections", {"values": {"Project ID": "UH-101", "Item": "Item", "Qty": "nan"}}).status_code, 400)

    def test_locked_workbook_keeps_original(self):
        before = self.path.read_bytes()
        with patch("workbook_store.os.replace", side_effect=PermissionError()):
            response = self.post("/api/projects", {"values": {"Project ID": "TEST-1", "Project Name": "Test House"}})
        self.assertEqual(response.status_code, 409)
        self.assertIn("Close", response.json["error"])
        self.assertEqual(self.path.read_bytes(), before)

    def test_csrf_and_host_protection(self):
        self.assertEqual(self.client.post("/api/projects", json={}).status_code, 403)
        self.assertEqual(self.client.get("/", headers={"Host": "attacker.example"}).status_code, 400)
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("frame-ancestors 'none'", response.headers["Content-Security-Policy"])
        for name in ["ui.css", "ui.js"]:
            with self.client.get(f"/assets/{name}") as asset:
                self.assertEqual(asset.status_code, 200)
        for name in ["symbol", "wordmark", "icon"]:
            with self.client.get(f"/brand/{name}.png") as asset:
                self.assertEqual(asset.status_code, 200)
                self.assertEqual(asset.mimetype, "image/png")
                self.assertTrue(asset.data.startswith(b"\x89PNG"))
        self.assertEqual(self.client.get("/brand/unknown.png").status_code, 404)


class RemoteTests(unittest.TestCase):
    def test_private_and_non_http_urls_are_rejected(self):
        import remote
        for url in ["file:///etc/passwd", "javascript:alert(1)", "https://user:password@example.com", "http://example.com:8080"]:
            with self.assertRaises(ValueError):
                remote.public_url(url)
        with patch("remote.socket.getaddrinfo", return_value=[(2, 1, 6, "", ("127.0.0.1", 80))]):
            with self.assertRaises(ValueError):
                remote.public_url("http://example.com")


class BrandingTests(unittest.TestCase):
    def test_logos_retain_transparency_and_aspect_ratio(self):
        from branding import load_logo
        from build_lookbook import CFG, logo
        from unittest.mock import Mock
        for field in ["logo_path", "logo_mark_path"]:
            image = load_logo(CFG[field])
            self.assertIsNotNone(image)
            self.assertEqual(image.mode, "RGBA")
            self.assertEqual(image.getchannel("A").getextrema(), (0, 255))
        canvas = Mock()
        self.assertTrue(logo(canvas, 396, 346, 40, max_width=300, centered=True))
        _, x, y, width, height = canvas.drawImage.call_args.args
        image = load_logo(CFG["logo_path"])
        self.assertAlmostEqual(width / height, image.width / image.height)
        self.assertAlmostEqual(x + width / 2, 396)
        self.assertLessEqual(width, 300)
        self.assertEqual(canvas.drawImage.call_args.kwargs["mask"], "auto")

    def test_branding_embedded_in_pdf_and_powerpoint(self):
        import build_lookbook
        from PIL import Image
        from pptx import Presentation
        from workbook_store import SELECTION_FIELDS
        project = {"Project ID": "BRAND", "Project Name": "Branding preview"}
        row = {**{key: "" for key in SELECTION_FIELDS}, "Section": "Kitchen", "Item": "Test fixture"}
        with tempfile.TemporaryDirectory() as folder:
            pdf = build_lookbook.build("BRAND", project, [row], True, Path(folder))
            pptx = build_powerpoint.build("BRAND", project, [row], True, Path(folder))
            self.assertIn(b"/SMask", pdf.read_bytes())
            prs = Presentation(pptx)
            for slide in [prs.slides[0], prs.slides[-1]]:
                pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
                logos = [shape for shape in pictures if shape.image.size == (894, 132)]
                self.assertEqual(len(logos), 1)
                with Image.open(io.BytesIO(logos[0].image.blob)) as image:
                    self.assertEqual(image.mode, "RGBA")
                    self.assertEqual(image.getchannel("A").getextrema(), (0, 255))
                self.assertAlmostEqual(logos[0].width / logos[0].height, 894 / 132, places=4)


class PresentationTests(unittest.TestCase):
    def test_local_images_are_embedded_in_both_formats(self):
        import build_lookbook
        import build_powerpoint
        from PIL import Image
        from pptx import Presentation
        from workbook_store import SELECTION_FIELDS
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "product.png"
            Image.new("RGBA", (120, 80), (100, 140, 110, 180)).save(image)
            project = {"Project ID": "IMAGE", "Project Name": "Image test", "Cover Image": str(image)}
            row = {**{key: "" for key in SELECTION_FIELDS}, "Section": "Kitchen", "Item": "Test fixture", "Image URL": str(image), "Product URL": "https://example.com/product", "Lookup Status": "Verified"}
            with patch("build_lookbook.CACHE", root / "cache"):
                pdf = build_lookbook.build("IMAGE", project, [row], False, root)
                pptx = build_powerpoint.build("IMAGE", project, [row], False, root)
            self.assertIn(b"/Subtype /Image", pdf.read_bytes())
            self.assertIn(b"https://example.com/product", pdf.read_bytes())
            prs = Presentation(pptx)
            self.assertTrue(any(shape.shape_type == 13 and shape.image.size == (120, 80) for slide in prs.slides for shape in slide.shapes))
            self.assertTrue(any(run.hyperlink.address == row["Product URL"] for slide in prs.slides for shape in slide.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs for run in p.runs))

    def test_many_sections_paginate_pdf_and_pptx(self):
        import build_lookbook
        import build_powerpoint
        from workbook_store import SELECTION_FIELDS
        project = {"Project ID": "TEST", "Project Name": "A test home with many categories", "Client Name": "Test client"}
        rows = [{**{key: "" for key in SELECTION_FIELDS}, "Section": f"Section {i:02}", "Item": f"Item {i}", "Manufacturer": "Test", "Lookup Status": "Not run"} for i in range(31)]
        with tempfile.TemporaryDirectory() as folder, patch("build_lookbook.get_image", return_value=None), patch("build_powerpoint.get_image", return_value=None):
            pdf = build_lookbook.build("TEST", project, rows, True, Path(folder))
            pptx = build_powerpoint.build("TEST", project, rows, True, Path(folder))
            self.assertTrue(pdf.read_bytes().startswith(b"%PDF-"))
            self.assertTrue(pptx.read_bytes().startswith(b"PK"))


if __name__ == "__main__":
    unittest.main()
