"""Hosted deployment: the Supabase-backed workbook, sign-in, and exposure guards.

No network is used. FakeSupabase stands in for Storage and PostgREST, implementing just
enough of their filter semantics to exercise the compare-and-swap that settles saves.
"""
import hashlib
import io
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import common
from openpyxl import load_workbook
from workbook_store import ConflictError, ValidationError, WorkbookStore, put, records

import auth
import deploy_render
import supabase_store
from supabase_store import HISTORY, LOOKUP_STATE, STATE, VERSIONS, SupabaseBackend

CONFIG = {"url": "https://project.supabase.co", "key": "service-key-for-tests",
          "bucket": "workbooks", "name": "UH_Homes_Selections_Tracker.xlsx"}


class FakeSupabase:
    """In-memory Storage plus the slice of PostgREST the backend relies on."""

    def __init__(self):
        self.objects = {}
        self.tables = {STATE: [], VERSIONS: [], HISTORY: [], LOOKUP_STATE: []}
        self.selects = {}
        self.downloads = 0
        self.uploads = 0

    # --- Storage ---
    def download(self, path):
        if path not in self.objects:
            raise RuntimeError("Supabase could not read the stored workbook (HTTP 404).")
        self.downloads += 1
        return self.objects[path]

    def upload(self, path, raw):
        self.uploads += 1
        self.objects[path] = raw

    def remove(self, path):
        self.objects.pop(path, None)

    # --- PostgREST ---
    def _match(self, rows, params):
        for key, value in params.items():
            if key in {"select", "order", "limit", "offset"}:
                continue
            op, _, wanted = str(value).partition(".")
            if op == "eq":
                rows = [r for r in rows if str(r.get(key)) == wanted]
            elif op == "neq":
                rows = [r for r in rows if str(r.get(key, "")) != wanted]
            elif op == "like":
                prefix = wanted.replace(r"\%", "%").replace(r"\_", "_").rstrip("*")
                rows = [r for r in rows if str(r.get(key, "")).startswith(prefix)]
            else:
                raise AssertionError(f"FakeSupabase does not implement operator {op!r}")
        return rows

    def select(self, table, params, count=False):
        self.selects[table] = self.selects.get(table, 0) + 1
        rows = self._match(list(self.tables[table]), params)
        if "order" in params:
            column, _, direction = params["order"].partition(".")
            rows.sort(key=lambda r: str(r.get(column) or ""), reverse=direction == "desc")
        if count:
            return len(rows)
        rows = rows[int(params.get("offset", 0)):]
        return rows[:int(params["limit"])] if "limit" in params else rows

    def _key(self, table):
        return {STATE: "id", VERSIONS: "revision", LOOKUP_STATE: "key", HISTORY: None}[table]

    def insert(self, table, row, upsert=False):
        key = self._key(table)
        if key:
            existing = next((r for r in self.tables[table] if str(r.get(key)) == str(row.get(key))), None)
            if existing and not upsert:
                raise RuntimeError(f"Supabase could not write {table} (HTTP 409).")
            if existing:
                existing.update(row)
                return [existing]
        row = {"created_at": f"2025-10-07T00:00:{len(self.tables[table]):02d}", **row}
        self.tables[table].append(row)
        return [row]

    def update(self, table, params, row):
        matched = self._match(self.tables[table], params)
        for existing in matched:
            existing.update(row)
        return matched


def hosted_backend(workbook):
    api = FakeSupabase()
    backend = SupabaseBackend(CONFIG, api=api)
    raw = Path(workbook).read_bytes()
    revision = hashlib.sha256(raw).hexdigest()
    path = "versions/seed.xlsx"
    api.upload(path, raw)
    api.insert(VERSIONS, {"revision": revision, "name": "UH_Homes_Selections_Tracker_20251007_000000_0.xlsx",
                          "object_path": path, "size": len(raw)})
    api.insert(STATE, {"id": 1, "revision": revision, "object_path": path})
    return backend, api


class SupabaseBackendTests(unittest.TestCase):
    def setUp(self):
        self.backend, self.api = hosted_backend(common.WORKBOOK)
        self.store = WorkbookStore(self.backend)

    def test_reads_the_object_named_by_the_pointer_row(self):
        raw, revision = self.backend.read()
        self.assertEqual(revision, self.api.tables[STATE][0]["revision"])
        self.assertEqual(hashlib.sha256(raw).hexdigest(), revision)

    def test_unchanged_revision_is_served_from_cache(self):
        """Versions are immutable, so one download per revision is always enough."""
        self.backend.read()
        self.backend.read()
        self.assertEqual(self.api.downloads, 1)

    def test_save_round_trips_through_storage(self):
        wb, revision = self.store.snapshot()
        sheet = common.Sheet(wb["Selections"])
        row = next(r for r, _ in sheet.rows())
        sheet.set(row, "Client Notes", "Hosted edit")
        self.store.save(wb, revision, "hosted test")
        stored = records(self.store.snapshot()[0], "Selections")
        self.assertEqual(next(r for r in stored if r["_row"] == row)["Client Notes"], "Hosted edit")

    def test_save_advances_the_pointer_and_keeps_the_previous_object(self):
        wb, revision = self.store.snapshot()
        common.Sheet(wb["Selections"]).set(2, "Client Notes", "Moved on")
        self.store.save(wb, revision)
        self.assertNotEqual(self.api.tables[STATE][0]["revision"], revision)
        self.assertEqual(len(self.api.objects), 2)          # the earlier version is the saved copy
        self.assertEqual(len(self.api.tables[VERSIONS]), 2)

    def test_a_second_writer_loses_the_compare_and_swap(self):
        """Two processes holding the same revision must not both succeed."""
        first, revision = self.store.snapshot()
        second, _ = self.store.snapshot()
        common.Sheet(first["Selections"]).set(2, "Client Notes", "First writer")
        self.store.save(first, revision)
        common.Sheet(second["Selections"]).set(2, "Client Notes", "Second writer")
        with self.assertRaises(ConflictError):
            self.store.save(second, revision)
        self.assertEqual(records(self.store.snapshot()[0], "Selections")[0]["Client Notes"], "First writer")

    def test_a_lost_save_leaves_no_orphan_object_or_copy(self):
        """The loser of a race must leave nothing behind that could later be restored."""
        other = WorkbookStore(SupabaseBackend(CONFIG, api=self.api))
        other.backend._ttl = 3600
        stale, stale_revision = other.snapshot()

        wb, revision = self.store.snapshot()
        put(common.Sheet(wb["Selections"]), 2, {"Client Notes": "Winner"})
        self.store.save(wb, revision)
        objects, copies = set(self.api.objects), self.store.backups()

        put(common.Sheet(stale["Selections"]), 2, {"Client Notes": "Loser"})
        with self.assertRaises(ConflictError):
            other.save(stale, stale_revision)
        self.assertEqual(set(self.api.objects), objects)
        # A version row whose object was never reached must not be offered as a restorable copy.
        self.assertEqual(self.store.backups(), copies)

    def test_every_save_is_a_distinct_revision(self):
        """openpyxl stamps the write time into the file, so no two saves can collide."""
        wb, revision = self.store.snapshot()
        row = next(r for r, _ in common.Sheet(wb["Selections"]).rows())
        put(common.Sheet(wb["Selections"]), row, {"Client Notes": "Temporary"})
        self.store.save(wb, revision)
        wb, revision = self.store.snapshot()
        put(common.Sheet(wb["Selections"]), row, {"Client Notes": ""})
        self.store.save(wb, revision)
        revisions = [r["revision"] for r in self.api.tables[VERSIONS]]
        self.assertEqual(len(set(revisions)), 3)
        self.assertEqual(records(self.store.snapshot()[0], "Selections")[0]["Client Notes"], "")

    def test_backups_list_and_restore_use_the_version_history(self):
        wb, revision = self.store.snapshot()
        row = next(r for r, _ in common.Sheet(wb["Selections"]).rows())
        common.Sheet(wb["Selections"]).set(row, "Client Notes", "First note")
        self.store.save(wb, revision)
        wb, revision = self.store.snapshot()
        common.Sheet(wb["Selections"]).set(row, "Client Notes", "Second note")
        self.store.save(wb, revision)

        copies = self.store.backups()                       # earlier versions, newest first
        self.assertEqual(len(copies), 2)
        notes = []
        for copy in copies:
            restored = load_workbook(io.BytesIO(self.store.read_backup(copy["name"])), keep_links=False)
            sheet = common.Sheet(restored["Selections"])
            notes.append(common.clean(restored["Selections"].cell(row, sheet.cols["Client Notes"]).value))
        self.assertEqual(notes, ["First note", ""])         # the live "Second note" is not a copy

    def test_restoring_an_unknown_copy_is_rejected(self):
        with self.assertRaises(ValidationError):
            self.store.read_backup("UH_Homes_Selections_Tracker_20200101_000000_0.xlsx")

    def test_export_revisions_are_counted_per_project(self):
        self.store.log("export UH-101 pdf final")
        self.store.log("export UH-101 pptx final")
        self.store.log("export UH-202 pdf final")
        self.assertEqual(self.store.count_notes("export UH-101 "), 2)

    def test_like_wildcards_in_a_note_prefix_are_escaped(self):
        self.store.log("export 100%_real pdf")
        self.store.log("export 100XXreal pdf")
        self.assertEqual(self.store.count_notes("export 100%_real "), 1)

    def test_the_pointer_row_is_held_briefly_across_polls(self):
        """The interface polls every few seconds; that must not be one round trip each time."""
        self.backend.read()
        before = self.api.selects[STATE]
        self.backend.read()
        self.backend.read()
        self.assertEqual(self.api.selects[STATE], before)

    def test_a_stale_cached_pointer_still_cannot_overwrite_another_writer(self):
        """The cache may lag, but the compare-and-swap, not the cache, decides who wins."""
        other = WorkbookStore(SupabaseBackend(CONFIG, api=self.api))
        other.backend._ttl = 3600                       # this process never notices the other's save
        stale, stale_revision = other.snapshot()

        wb, revision = self.store.snapshot()
        common.Sheet(wb["Selections"]).set(2, "Client Notes", "Winner")
        self.store.save(wb, revision)

        common.Sheet(stale["Selections"]).set(2, "Client Notes", "Stale writer")
        with self.assertRaises(ConflictError):
            other.save(stale, stale_revision)
        self.assertEqual(self.api.tables[STATE][0]["revision"], hashlib.sha256(self.backend.read()[0]).hexdigest())

    def test_the_hosted_workbook_is_never_locked_by_excel(self):
        self.assertFalse(self.store.locked_elsewhere())

    def test_lookup_state_survives_a_round_trip(self):
        state = self.store.state()
        self.assertIsNone(state.read())
        state.write('{"entries": {}, "pending": {}}')
        self.assertEqual(state.read(), '{"entries": {}, "pending": {}}')

    def test_an_empty_pointer_table_explains_how_to_seed(self):
        self.api.tables[STATE].clear()
        with self.assertRaises(ValidationError) as caught:
            self.store.snapshot()
        self.assertIn("push", str(caught.exception))

    def test_the_hosted_workbook_has_no_local_path(self):
        with self.assertRaises(AttributeError):
            self.store.path


class SettingsTests(unittest.TestCase):
    def test_no_configuration_means_local(self):
        with patch.dict("os.environ", {"SUPABASE_URL": "", "SUPABASE_SERVICE_KEY": ""}):
            self.assertIsNone(supabase_store.settings())

    def test_half_configured_counts_as_local_rather_than_breaking_the_app(self):
        """Either half alone is a setup in progress, not a fatal error."""
        for half in [{"SUPABASE_URL": "https://p.supabase.co", "SUPABASE_SERVICE_KEY": ""},
                     {"SUPABASE_URL": "", "SUPABASE_SERVICE_KEY": "sb_secret_x"}]:
            with patch.dict("os.environ", half):
                self.assertIsNone(supabase_store.settings())

    def test_a_url_with_a_path_is_refused(self):
        with patch.dict("os.environ", {"SUPABASE_URL": "https://p.supabase.co/rest/v1",
                                       "SUPABASE_SERVICE_KEY": "key"}):
            with self.assertRaises(supabase_store.ConfigError):
                supabase_store.settings()

    def test_keys_never_appear_in_an_upload_failure(self):
        api = supabase_store.Supabase(CONFIG)

        class Response:
            status_code = 500
            text = "service-key-for-tests leaked"

        with self.assertRaises(RuntimeError) as caught:
            api._check(Response(), "store the workbook")
        self.assertNotIn(CONFIG["key"], str(caught.exception))


class ApiKeyHeaderTests(unittest.TestCase):
    """Supabase's newer keys are opaque strings, not JWTs, and only belong on `apikey`."""

    def test_a_new_style_secret_key_is_not_sent_as_a_bearer_token(self):
        headers = auth.key_headers("sb_secret_4TRCaexample")
        self.assertEqual(headers, {"apikey": "sb_secret_4TRCaexample"})

    def test_a_new_style_publishable_key_is_not_sent_as_a_bearer_token(self):
        self.assertNotIn("Authorization", auth.key_headers("sb_publishable_J97yXQexample"))

    def test_a_legacy_jwt_key_is_sent_on_both_headers(self):
        legacy = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.payload.signature"
        self.assertEqual(auth.key_headers(legacy),
                         {"apikey": legacy, "Authorization": f"Bearer {legacy}"})

    def test_the_client_applies_the_rule_and_looks_nothing_like_a_browser(self):
        api = supabase_store.Supabase({**CONFIG, "key": "sb_secret_4TRCaexample"})
        self.assertEqual(api.session.headers["apikey"], "sb_secret_4TRCaexample")
        self.assertNotIn("Authorization", api.session.headers)
        # A browser-like User-Agent would have the secret key rejected with 401.
        self.assertNotIn("Mozilla", api.session.headers["User-Agent"])


class SignInTests(unittest.TestCase):
    CONFIG = {"url": "https://project.supabase.co", "key": "anon-key"}

    class Client:
        def __init__(self, status, body=None):
            self.status, self.body, self.calls = status, body or {}, 0

        def post(self, *args, **kwargs):
            self.calls += 1
            client = self

            class Response:
                status_code = client.status

                def json(self):
                    return client.body
            return Response()

    def test_a_bad_address_is_refused_without_a_request(self):
        client = self.Client(200)
        with self.assertRaises(auth.AuthError):
            auth.sign_in("not-an-email", "password123", self.CONFIG, client)
        self.assertEqual(client.calls, 0)

    def test_a_short_password_is_refused_without_a_request(self):
        client = self.Client(200)
        with self.assertRaises(auth.AuthError):
            auth.sign_in("team@uh.example", "x", self.CONFIG, client)
        self.assertEqual(client.calls, 0)

    def test_rejection_does_not_reveal_whether_the_account_exists(self):
        with self.assertRaises(auth.AuthError) as caught:
            auth.sign_in("team@uh.example", "wrong-password", self.CONFIG, self.Client(400))
        self.assertIn("do not match an account", str(caught.exception))

    def test_a_successful_sign_in_returns_the_identity_only(self):
        body = {"access_token": "secret-token", "user": {"id": "abc", "email": "team@uh.example"}}
        user = auth.sign_in("team@uh.example", "password123", self.CONFIG, self.Client(200, body))
        self.assertEqual(user, {"id": "abc", "email": "team@uh.example"})

    def test_creating_an_account_generates_a_password_and_confirms_it(self):
        sent = {}

        class Client:
            def post(self, url, json=None, headers=None, timeout=None):
                sent.update(url=url, body=json)
                return type("R", (), {"status_code": 200, "json": lambda self: {"id": "abc"}})()

        with patch.dict("os.environ", {"SUPABASE_URL": "https://p.supabase.co",
                                       "SUPABASE_SERVICE_KEY": "service"}):
            created = auth.create_user("team@uh.example", session=Client())
        self.assertTrue(sent["url"].endswith("/auth/v1/admin/users"))
        self.assertTrue(sent["body"]["email_confirm"])          # no email delivery needed
        self.assertGreaterEqual(len(created["password"]), 16)
        self.assertEqual(created["email"], "team@uh.example")

    def test_creating_an_account_without_configuration_is_refused(self):
        with patch.dict("os.environ", {"SUPABASE_URL": "", "SUPABASE_SERVICE_KEY": ""}):
            with self.assertRaises(auth.AuthError):
                auth.create_user("team@uh.example")

    def test_a_duplicate_account_is_reported_clearly(self):
        class Client:
            def post(self, *args, **kwargs):
                return type("R", (), {"status_code": 422,
                                      "json": lambda self: {"msg": "User already registered"}})()

        with patch.dict("os.environ", {"SUPABASE_URL": "https://p.supabase.co",
                                       "SUPABASE_SERVICE_KEY": "service"}):
            with self.assertRaises(auth.AuthError) as caught:
                auth.create_user("team@uh.example", session=Client())
        self.assertIn("already exists", str(caught.exception))

    def test_repeated_failures_are_throttled(self):
        throttle = auth.Throttle(limit=3, window=60)
        for _ in range(3):
            throttle.check("1.2.3.4")
            throttle.record("1.2.3.4")
        with self.assertRaises(auth.AuthError):
            throttle.check("1.2.3.4")
        throttle.check("5.6.7.8")       # another address is unaffected


class RenderDeployTests(unittest.TestCase):
    """Creating the host without hand-copying settings or guessing the hostname."""

    ENV = {"SUPABASE_URL": "https://p.supabase.co", "SUPABASE_SERVICE_KEY": "service",
           "SUPABASE_ANON_KEY": "anon", "UH_SECRET_KEY": "x" * 64, "RENDER_API_KEY": "rnd_key",
           "SEARCH_PROVIDER": "serpapi", "SERPAPI_API_KEY": "serp", "BRAVE_API_KEY": "",
           "SERPAPI_KEY": "", "ANTHROPIC_API_KEY": "", "UH_TRUSTED_HOSTS": "", "SUPABASE_BUCKET": ""}

    class Session:
        """Records calls and replays canned Render responses."""

        def __init__(self, existing=None, status=200):
            self.headers = {}
            self.calls = []
            self.existing = existing
            self.status = status

        def request(self, method, url, timeout=None, json=None, params=None):
            self.calls.append({"method": method, "url": url, "json": json, "params": params})
            session = self

            class Response:
                status_code = session.status
                content = b"{}"

                def json(self):
                    if url.endswith("/owners"):
                        return [{"owner": {"id": "usr-1", "name": "Sai", "email": "s@uh.example",
                                           "type": "user"}}]
                    if url.endswith("/services") and method == "GET":
                        return session.existing or []
                    if url.endswith("/services") and method == "POST":
                        return {"service": {"id": "srv-9", "name": "uh-selections", "slug": "uh-selections",
                                            "serviceDetails": {"url": "https://uh-selections-ab12.onrender.com"}}}
                    return {}
            return Response()

    def test_a_new_service_is_created_and_configured(self):
        session = self.Session()
        with patch.dict("os.environ", self.ENV):
            result = deploy_render.deploy(session=session)
        self.assertTrue(result["created"])
        created = next(c for c in session.calls if c["method"] == "POST")
        details = created["json"]["serviceDetails"]
        self.assertEqual(details["runtime"], "python")
        self.assertEqual(details["plan"], "free")
        self.assertEqual(details["healthCheckPath"], "/healthz")
        self.assertEqual(details["numInstances"], 1)      # one instance, never more
        self.assertIn("waitress-serve", details["envSpecificDetails"]["startCommand"])

    def test_trusted_hosts_is_taken_from_the_hostname_render_assigned(self):
        """The hostname is unknown until the service exists, and a wrong value 400s every request."""
        session = self.Session()
        with patch.dict("os.environ", self.ENV):
            result = deploy_render.deploy(session=session)
        self.assertEqual(result["hostname"], "uh-selections-ab12.onrender.com")
        sent = next(c for c in session.calls if c["method"] == "PUT")["json"]
        variables = {entry["key"]: entry["value"] for entry in sent}
        self.assertEqual(variables["UH_TRUSTED_HOSTS"], "uh-selections-ab12.onrender.com")
        self.assertEqual(variables["UH_HOST"], "0.0.0.0")

    def test_empty_local_settings_are_not_sent_as_blanks(self):
        session = self.Session()
        with patch.dict("os.environ", self.ENV):
            deploy_render.deploy(session=session)
        variables = {e["key"]: e["value"] for e in next(c for c in session.calls if c["method"] == "PUT")["json"]}
        self.assertIn("SERPAPI_API_KEY", variables)
        for absent in ["BRAVE_API_KEY", "SERPAPI_KEY", "ANTHROPIC_API_KEY", "SUPABASE_BUCKET"]:
            self.assertNotIn(absent, variables)

    def test_an_existing_service_is_reused_rather_than_duplicated(self):
        existing = [{"service": {"id": "srv-old", "name": "uh-selections", "slug": "uh-selections",
                                 "serviceDetails": {"url": "https://uh-selections.onrender.com"}}}]
        session = self.Session(existing=existing)
        with patch.dict("os.environ", self.ENV):
            result = deploy_render.deploy(session=session)
        self.assertFalse(result["created"])
        self.assertEqual(result["service_id"], "srv-old")
        self.assertFalse([c for c in session.calls if c["method"] == "POST"])

    def test_missing_supabase_settings_stop_the_deploy_before_it_configures_anything(self):
        session = self.Session()
        with patch.dict("os.environ", {**self.ENV, "SUPABASE_SERVICE_KEY": ""}):
            with self.assertRaises(deploy_render.DeployError) as caught:
                deploy_render.deploy(session=session)
        self.assertIn("SUPABASE_SERVICE_KEY", str(caught.exception))
        self.assertFalse([c for c in session.calls if c["method"] == "PUT"])

    def test_a_missing_api_key_is_explained(self):
        with patch.dict("os.environ", {**self.ENV, "RENDER_API_KEY": ""}):
            with self.assertRaises(deploy_render.DeployError) as caught:
                deploy_render.deploy(session=self.Session())
        self.assertIn("RENDER_API_KEY", str(caught.exception))

    def test_a_rejected_key_is_explained(self):
        with patch.dict("os.environ", self.ENV):
            with self.assertRaises(deploy_render.DeployError) as caught:
                deploy_render.deploy(session=self.Session(status=401))
        self.assertIn("rejected the API key", str(caught.exception))

    def test_several_workspaces_ask_which_one(self):
        session = self.Session()

        def two(method, url, timeout=None, json=None, params=None):
            class Response:
                status_code = 200
                content = b"{}"

                def json(self):
                    return [{"owner": {"id": "usr-1", "name": "Personal", "email": "a@x.example", "type": "user"}},
                            {"owner": {"id": "tea-2", "name": "Studio", "email": "b@x.example", "type": "team"}}]
            return Response()

        session.request = two
        with patch.dict("os.environ", self.ENV):
            with self.assertRaises(deploy_render.DeployError) as caught:
                deploy_render.deploy(session=session)
        self.assertIn("--owner-email", str(caught.exception))


class HostedAppTests(unittest.TestCase):
    """The guard that keeps client selections and pricing behind sign-in."""

    def setUp(self):
        import app
        self.module = app
        self.folder = tempfile.TemporaryDirectory()
        self.path = Path(self.folder.name) / "master.xlsx"
        shutil.copy2(common.WORKBOOK, self.path)
        self.patches = [patch.object(app, "store", WorkbookStore(self.path)),
                        patch.object(app, "automation", None),
                        patch.object(app, "REQUIRE_LOGIN", True),
                        patch.object(app, "AUTH", SignInTests.CONFIG)]
        for item in self.patches:
            item.start()
        app.app.config["SECRET_KEY"] = "test-secret"
        self.client = app.app.test_client()

    def tearDown(self):
        for item in reversed(self.patches):
            item.stop()
        self.folder.cleanup()

    def test_api_requires_a_session(self):
        response = self.client.get("/api/state")
        self.assertEqual(response.status_code, 401)
        self.assertIn("Sign in", response.json["error"])

    def test_every_mutating_route_requires_a_session(self):
        for path in ["/api/projects", "/api/selections", "/api/presentation", "/api/workbook/import",
                     "/api/backups/restore", "/api/automation/retry"]:
            self.assertEqual(self.client.post(path, json={}).status_code, 401, path)

    def test_the_workbook_download_requires_a_session(self):
        self.assertEqual(self.client.get("/api/workbook").status_code, 401)
        self.assertEqual(self.client.get("/api/backups").status_code, 401)
        self.assertEqual(self.client.get("/api/selections/csv?project=UH-101").status_code, 401)

    def test_the_interface_redirects_to_sign_in(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/login", response.headers["Location"])

    def test_the_health_check_stays_open(self):
        self.assertEqual(self.client.get("/healthz").status_code, 200)

    def test_the_health_check_answers_the_platforms_internal_hostname(self):
        # Render's liveness probe arrives on an internal host; TRUSTED_HOSTS must
        # not reject it or the deploy times out as "unhealthy".
        response = self.client.get("/healthz", headers={"Host": "srv-abc123.internal"})
        self.assertEqual(response.status_code, 200)

    def test_other_routes_still_reject_an_unknown_host(self):
        response = self.client.get("/", headers={"Host": "not-our-site.example"})
        self.assertEqual(response.status_code, 400)

    def test_the_sign_in_page_renders_without_a_session(self):
        response = self.client.get("/login")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Sign in", response.data)

    def test_signing_in_opens_the_workspace(self):
        page = self.client.get("/login")
        token = page.data.split(b'name="token" value="')[1].split(b'"')[0].decode()
        body = {"user": {"id": "abc", "email": "team@uh.example"}}
        with patch("auth.requests.post", return_value=type("R", (), {"status_code": 200, "json": lambda self: body})()):
            response = self.client.post("/login", data={"token": token, "email": "team@uh.example",
                                                        "password": "password123", "next": "/"})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.client.get("/api/state").status_code, 200)

    def test_a_sign_in_without_the_form_token_is_refused(self):
        self.client.get("/login")
        response = self.client.post("/login", data={"token": "wrong", "email": "team@uh.example",
                                                    "password": "password123"})
        self.assertEqual(response.status_code, 400)

    def test_the_session_token_guards_writes_after_signing_in(self):
        page = self.client.get("/login")
        token = page.data.split(b'name="token" value="')[1].split(b'"')[0].decode()
        body = {"user": {"id": "abc", "email": "team@uh.example"}}
        with patch("auth.requests.post", return_value=type("R", (), {"status_code": 200, "json": lambda self: body})()):
            self.client.post("/login", data={"token": token, "email": "team@uh.example",
                                             "password": "password123", "next": "/"})
        revision = self.client.get("/api/state").json["revision"]
        values = {"Project ID": "TEST-1", "Project Name": "Test House"}
        # The pre-login form token is rotated by signing in, so it must no longer be accepted.
        stale = self.client.post("/api/projects", json={"revision": revision, "create": True, "values": values},
                                 headers={"X-UH-Token": token})
        self.assertEqual(stale.status_code, 403)
        fresh = self.client.get("/").data.split(b'name="uh-token" content="')[1].split(b'"')[0].decode()
        response = self.client.post("/api/projects", json={"revision": revision, "create": True, "values": values},
                                    headers={"X-UH-Token": fresh})
        self.assertEqual(response.status_code, 200, response.json)

    def test_signing_out_ends_the_session(self):
        page = self.client.get("/login")
        token = page.data.split(b'name="token" value="')[1].split(b'"')[0].decode()
        body = {"user": {"id": "abc", "email": "team@uh.example"}}
        with patch("auth.requests.post", return_value=type("R", (), {"status_code": 200, "json": lambda self: body})()):
            self.client.post("/login", data={"token": token, "email": "team@uh.example",
                                             "password": "password123", "next": "/"})
        fresh = self.client.get("/").data.split(b'name="uh-token" content="')[1].split(b'"')[0].decode()
        self.assertEqual(self.client.post("/api/logout", headers={"X-UH-Token": fresh}).status_code, 200)
        self.assertEqual(self.client.get("/api/state").status_code, 401)


class ExposureTests(unittest.TestCase):
    """A deployment reachable by others must not start without sign-in."""

    def test_a_local_bind_needs_no_sign_in(self):
        import app
        with patch.object(app, "LOCAL_ONLY", True), patch.object(app, "REQUIRE_LOGIN", False):
            app.check_exposure()

    def test_a_public_bind_without_sign_in_refuses_to_start(self):
        import app
        with patch.object(app, "LOCAL_ONLY", False), patch.object(app, "REQUIRE_LOGIN", False):
            with self.assertRaises(SystemExit) as caught:
                app.check_exposure()
        self.assertIn("sign-in is not configured", str(caught.exception))

    def test_a_public_bind_needs_a_persistent_secret_key(self):
        import app
        with patch.object(app, "LOCAL_ONLY", False), patch.object(app, "REQUIRE_LOGIN", True), \
             patch.dict("os.environ", {"UH_SECRET_KEY": "", "UH_TRUSTED_HOSTS": "uh.example.com"}):
            with self.assertRaises(SystemExit) as caught:
                app.check_exposure()
        self.assertIn("UH_SECRET_KEY", str(caught.exception))

    def test_a_public_bind_needs_trusted_hosts(self):
        import app
        with patch.object(app, "LOCAL_ONLY", False), patch.object(app, "REQUIRE_LOGIN", True), \
             patch.dict("os.environ", {"UH_SECRET_KEY": "x" * 64, "UH_TRUSTED_HOSTS": ""}):
            with self.assertRaises(SystemExit) as caught:
                app.check_exposure()
        self.assertIn("UH_TRUSTED_HOSTS", str(caught.exception))

    def test_a_public_bind_refuses_to_serve_a_local_workbook_file(self):
        """A host's filesystem is temporary, so this would silently lose every edit."""
        import app
        local = type("B", (), {"remote": False})()
        with patch.object(app, "LOCAL_ONLY", False), patch.object(app, "REQUIRE_LOGIN", True), \
             patch.object(app.store, "backend", local), \
             patch.dict("os.environ", {"UH_SECRET_KEY": "x" * 64, "UH_TRUSTED_HOSTS": "uh.example.com"}):
            with self.assertRaises(SystemExit) as caught:
                app.check_exposure()
        self.assertIn("redeploy would discard", str(caught.exception))

    def test_a_half_configured_env_still_runs_locally(self):
        """Mid-setup, with a URL but no key yet, the local app must keep working."""
        with patch.dict("os.environ", {"SUPABASE_URL": "https://p.supabase.co",
                                       "SUPABASE_SERVICE_KEY": ""}):
            self.assertIsNone(supabase_store.settings())
            reported = dict((message, ok) for ok, message in supabase_store.check())
        self.assertTrue(any("SUPABASE_SERVICE_KEY is empty" in m for m in reported))


if __name__ == "__main__":
    unittest.main()
